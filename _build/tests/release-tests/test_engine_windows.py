import sys,os,json,subprocess,time,hashlib,threading,queue
from pathlib import Path
ROOT=Path(__file__).resolve().parent;sys.path[:0]=[str(ROOT/'src'),'E:/四維八德/_build/vendor'];import engine as E
BASE='E:/四維八德';E.base_dir=lambda:BASE;E.FFMPEG=E.find_tool('ffmpeg');E.PANDOC=E.find_tool('pandoc');E.SOFFICE=E.find_soffice();E.CALIBRE=E._calibre();E.SEVENZIP=E.find_tool('7z');E.detect_hw=lambda:[]
DATA=ROOT/'windows-engine-checks';DATA.mkdir(exist_ok=True);E._DATA=str(DATA);fixtures=DATA/'中文 素材';fixtures.mkdir(exist_ok=True)
def ff(*args):
 r=E.run([E.FFMPEG,'-hide_banner','-y','-v','error',*map(str,args)],timeout=90);assert not r.returncode,r.stderr
video=fixtures/'影片 測試.mp4';ff('-f','lavfi','-i','testsrc2=size=640x360:rate=24:duration=3','-f','lavfi','-i','sine=frequency=440:duration=3','-c:v','libx264','-preset','ultrafast','-c:a','aac',video)
audio=fixtures/'聲音.wav';ff('-f','lavfi','-i','sine=frequency=440:duration=2',audio)
big5=fixtures/'繁中字幕.srt';big5.write_bytes('1\n00:00:00,000 --> 00:00:01,000\n中文測試\n'.encode('cp950'))
md=fixtures/'文件.md';md.write_text('# Offline document\n\nNative C engine test.\n',encoding='utf-8')
from docx import Document
word=fixtures/'文件.docx';d=Document();d.add_paragraph('Offline Word test.');d.save(word)
from openpyxl import Workbook
sheet=fixtures/'表格.xlsx';w=Workbook();w.active.append(['Item','Amount']);w.active.append(['Apple',20]);w.save(sheet)
from pptx import Presentation
ppt=fixtures/'簡報.pptx';r=Presentation();s=r.slides.add_slide(r.slide_layouts[0]);s.shapes.title.text='Offline PowerPoint';r.save(ppt)
import zipfile
archive=fixtures/'中文壓縮檔.zip'
with zipfile.ZipFile(archive,'w') as z:z.write(md,'資料/文件.md')
hdr=fixtures/'HDR.mp4';ff('-f','lavfi','-i','testsrc2=size=320x180:rate=12:duration=1','-c:v','libx265','-preset','ultrafast','-pix_fmt','yuv420p10le','-x265-params','colorprim=bt2020:transfer=smpte2084:colormatrix=bt2020nc',hdr)
client=subprocess.Popen([str(ROOT/'shell/out/engine-test.exe'),BASE,str(DATA)],stdin=subprocess.PIPE,stdout=subprocess.PIPE,stderr=open(DATA/'native-stderr.log','w'),encoding='utf-8',creationflags=E.NOWIN)
client.stdin.write(json.dumps(dict(base=BASE,data=str(DATA)),ensure_ascii=False)+'\n');client.stdin.flush()
responses=queue.Queue()
def reader():
 for line in client.stdout:
  try:responses.put(json.loads(line))
  except Exception:responses.put(RuntimeError(line))
threading.Thread(target=reader,daemon=True).start()
def call(m,**args):
 client.stdin.write(json.dumps(dict(m=m,**args),ensure_ascii=False)+'\n');client.stdin.flush();r=responses.get(timeout=40)
 if isinstance(r,Exception):raise r
 return r
def wait(ids,timeout=120):
 end=time.monotonic()+timeout
 while time.monotonic()<end:
  result=[j for j in call('status') if j['id'] in ids]
  if len(result)==len(ids) and all(j['state'] not in ('queued','running') for j in result):return result
  time.sleep(.04)
 raise TimeoutError(ids)
eng=E.Engine(1,light=1);records=[]
cases=[(f'video_{t}',[video],t,{}) for t in ('mp4','mkv','webm','mov','avi','gif','mp3','m4a','wav','flac','ogg','opus')]
cases += [('trim',[video],'mp4',dict(trim_start=.4,trim_end=1.8)),('size_cap',[video],'mp4',dict(target_mb=.15)),('merge',[video,video],'mp4',dict(merge=True)),('hdr',[hdr],'mp4',{}),('big5_vtt',[big5],'vtt',{}),('big5_ass',[big5],'ass',{})]
cases += [(f'markdown_{t}',[md],t,{}) for t in ('docx','html','txt','rtf')]
cases += [('office_word_pdf',[word],'pdf',{}),('office_sheet_pdf',[sheet],'pdf',{}),('office_ppt_pdf',[ppt],'pdf',{}),('archive_7z',[archive],'7z',{}),('archive_extract',[archive],'extract',{}),('ebook',[word],'epub',{})]
try:
 for name,files,target,extra in cases:
  started=time.monotonic();record=dict(name=name,target=target)
  try:
   cdir=DATA/'C'/name;pdir=DATA/'Python'/name;cdir.mkdir(parents=True,exist_ok=True);pdir.mkdir(parents=True,exist_ok=True)
   opt=dict(preset='balanced',hwaccel=False,fastcopy=False,resource_mode='quiet',out_dir=str(cdir));opt.update(extra);merge=opt.pop('merge',False)
   ids=call('start',files=list(map(str,files)),target=target,opt=opt,merge=merge);c=wait(ids)[0];assert c['state']=='done',c
   opt['out_dir']=str(pdir);pids=eng.submit(list(map(str,files)),target,opt,merge=merge)
   end=time.monotonic()+120
   while eng.jobs[pids[0]].state in ('queued','running') and time.monotonic()<end:time.sleep(.04)
   p=eng.jobs[pids[0]];assert p.state=='done',p.info()
   def payload(path):
    path=Path(path)
    if path.is_file():return path.read_bytes()
    return b''.join(str(f.relative_to(path)).encode('utf-8')+f.read_bytes() for f in sorted(path.rglob('*')) if f.is_file())
   cdata=payload(c['output']);pdata=payload(p.output);cs=len(cdata);ps=len(pdata)
   record.update(ok=True,c_bytes=cs,python_bytes=ps,equal_size=cs==ps,equal_bytes=cdata==pdata,validation=c['validation'])
   if target in ('vtt','ass'):assert '中文測試' in Path(c['output']).read_text(encoding='utf-8-sig')
  except Exception as ex:record.update(ok=False,error=str(ex))
  record['seconds']=round(time.monotonic()-started,2);records.append(record);print(json.dumps(record,ensure_ascii=False),flush=True)
 # A real failed native job can be retried after its input becomes available.
 missing=fixtures/'稍後出現.mp4';missing.unlink(missing_ok=True);od=DATA/'retry';od.mkdir(exist_ok=True);opts=dict(preset='balanced',hwaccel=False,resource_mode='quiet',out_dir=str(od))
 ids=call('start',files=[str(missing)],target='mp3',opt=opts);failed=wait(ids)[0];assert failed['state']=='error';missing.write_bytes(video.read_bytes());time.sleep(.1);rid=call('retry',id=ids[0]);assert rid and rid!=ids[0];assert wait([rid])[0]['state']=='done';records.append(dict(name='error_retry',ok=True))
 # Cancel after the actual native process has begun producing progress.
 longvideo=fixtures/'長影片.mp4';ff('-stream_loop','29','-i',video,'-c','copy',longvideo);od=DATA/'cancel';od.mkdir(exist_ok=True);opts['out_dir']=str(od);opts['codec']='hevc';ids=call('start',files=[str(longvideo)],target='mp4',opt=opts)
 until=time.monotonic()+15
 while time.monotonic()<until:
  job=next(j for j in call('status') if j['id']==ids[0])
  if job['state']=='running' and job['progress']>0:break
  time.sleep(.03)
 call('cancel',id=ids[0]);assert wait(ids)[0]['state']=='cancelled';time.sleep(.2);records.append(dict(name='cancel_running',ok=True))
except Exception as ex:records.append(dict(name='queue_controls',ok=False,error=str(ex)))
finally:
 eng.shutdown();call('quit');client.stdin.close();client.wait(timeout=10);(DATA/'results.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
print('PASSED',sum(r['ok'] for r in records),'/',len(records));raise SystemExit(not records or not all(r['ok'] for r in records))
