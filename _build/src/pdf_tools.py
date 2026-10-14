# SPDX-License-Identifier: AGPL-3.0-only
# Copyright (c) 2026 Wel Wang
"""Offline PDF operations written for SiWeiBaDe. Never sends document data online."""
import os, io, json, re, math, shutil, tempfile, threading, base64, html, difflib, hashlib, time
from functools import wraps
import engine as E
import pdf_schema as PS

LOCK=threading.RLock()
fitz=None
def local(fn):
    @wraps(fn)
    def wrapped(*args,**kw):
        global fitz
        with LOCK:
            if fitz is None: import pymupdf; fitz=pymupdf
            return fn(*args,**kw)
    return wrapped

def asset(name): return os.path.join(E.base_dir(),'tools',name)
_OCR_DATA=None
def tessdata():
    with LOCK:return _tessdata()
def _tessdata():
    # Tesseract's Windows C runtime cannot open every Unicode data path.
    # Keep the shipped files intact and make a verified local cache when needed.
    global _OCR_DATA
    source=asset('pdf-ocr/tessdata')
    if source.isascii() or not os.path.isdir(source):return source
    if _OCR_DATA and os.path.isdir(_OCR_DATA):return _OCR_DATA
    def short_path(path):
        if path.isascii():return path
        if os.name=='nt':
            import ctypes
            buf=ctypes.create_unicode_buffer(32768)
            if ctypes.windll.kernel32.GetShortPathNameW(str(path),buf,len(buf)) and buf.value.isascii():return buf.value
        return None
    direct=short_path(source)
    if direct:return direct
    choices=[tempfile.gettempdir(),os.path.join(os.environ.get('PUBLIC',r'C:\Users\Public'),'Documents')]
    for parent in choices:
        try:
            parent=short_path(parent)
            if not parent:continue
            cache=os.path.join(parent,'SiWeiBaDe-OCR-'+hashlib.sha256(os.path.abspath(source).encode('utf-8')).hexdigest()[:12]);os.makedirs(cache,exist_ok=True)
            for name in os.listdir(source):
                if not name.endswith('.traineddata'):continue
                src=os.path.join(source,name);dst=os.path.join(cache,name)
                with open(src,'rb') as f:expected=hashlib.file_digest(f,'sha256').hexdigest()
                actual=''
                if os.path.isfile(dst):
                    with open(dst,'rb') as f:actual=hashlib.file_digest(f,'sha256').hexdigest()
                if actual!=expected:
                    shutil.copyfile(src,dst+'.writing');os.replace(dst+'.writing',dst)
            _OCR_DATA=cache;return cache
        except OSError:continue
    raise ValueError('OCR 無法建立可讀取的本機模型快取，請確認暫存資料夾有寫入權限')
def capabilities():
    # check the shipped folder directly: tessdata() may build a hashed copy of every language file (non-ASCII
    # install paths), which belongs to the first OCR job, not to opening the window
    return dict(ocr=os.path.isfile(os.path.join(asset('pdf-ocr/tessdata'),'eng.traineddata')), ai=os.path.isfile(asset('pdf-ai/model.gguf')) and os.path.isfile(asset('pdf-ai/llama-completion.exe')),
                pdfa=os.path.isfile(asset('ghostscript/bin/gswin64c.exe')), office=bool(E.SOFFICE), catalogue=PS.CATALOG, sections=PS.SECTIONS, offline=True)
def cancelled(j):
    if j.cancelled: raise InterruptedError
def output(j, ext='pdf', suffix=''):
    root=j.opt.get('out_dir') or os.path.dirname(j.files[0]); os.makedirs(root,exist_ok=True)
    name=os.path.splitext(os.path.basename(j.files[0]))[0]
    return E.unique_path(os.path.join(root,name+(suffix or '_'+PS.TOOLS[[t['id'] for t in PS.TOOLS].index(j.tool)]['name'].replace('/','_'))+'.'+ext),avoid=j.files)
def folder(j,suffix=''):
    path=output(j,'files',suffix)[:-6]; path=E.unique_path(path);os.makedirs(path);j.output=path;j.keep_output=True;return path
def open_pdf(path,pw=''):
    doc=fitz.open(path)
    if not doc.is_pdf or not doc.page_count: doc.close();raise ValueError('文件沒有可讀取的 PDF 頁面')
    if doc.needs_pass and not doc.authenticate(str(pw or '')): doc.close();raise ValueError('PDF 需要正確的文件密碼')
    return doc
def indices(spec,count,allow_repeat=False):
    if not str(spec or '').strip(): return list(range(count))
    result=[];seen=set()
    for part in str(spec).replace('，',',').split(','):
        part=part.strip();match=re.fullmatch(r'(\d+)(?:\s*-\s*(\d*))?',part)
        if not match: raise ValueError('頁碼格式錯誤，請輸入例如 1,3-5,8-')
        first=int(match[1]);last=int(match[2] or count) if '-' in part else first
        if not 1<=first<=last<=count: raise ValueError(f'頁碼超出範圍，文件共有 {count} 頁')
        for n in range(first-1,last):
            if allow_repeat or n not in seen: result.append(n);seen.add(n)
    if not result: raise ValueError('沒有選取頁面')
    return result
def rect(value,page,normalized=False):
    if isinstance(value,str): value=[float(v) for v in value.split(',')]
    if not isinstance(value,(list,tuple)) or len(value)!=4: raise ValueError('請選取有效範圍')
    x,y,w,h=[float(v) for v in value]
    if not all(math.isfinite(v) for v in (x,y,w,h)) or w<=0 or h<=0: raise ValueError('選取範圍無效')
    if normalized: x*=page.rect.width;y*=page.rect.height;w*=page.rect.width;h*=page.rect.height
    selected=fitz.Rect(x,y,x+w,y+h)&page.rect
    if selected.is_empty or selected.width<1 or selected.height<1: raise ValueError('選取範圍超出頁面')
    return selected
def color(value):
    value=str(value or '#000000')
    if not re.fullmatch(r'#[0-9a-fA-F]{6}',value): raise ValueError('顏色格式無效')
    return tuple(int(value[i:i+2],16)/255 for i in (1,3,5))
def safe_actions(value):
    if isinstance(value,str):
        if len(value)>4*1024*1024: raise ValueError('編輯內容過大')
        value=json.loads(value or '[]')
    if not isinstance(value,list) or len(value)>500: raise ValueError('一次最多 500 個編輯項目')
    if not all(isinstance(a,dict) for a in value): raise ValueError('編輯項目格式不符')
    return value
def selected_page(doc,a):
    n=int(a.get('page',1))-1
    if not 0<=n<len(doc): raise ValueError('編輯項目的頁碼超出文件範圍')
    page=doc[n]
    if page.rotation: page.remove_rotation()
    return page
def save(doc,path,j,preserve_password=True):
    cancelled(j)
    tmp=path+'.writing'
    opts=dict(garbage=4,deflate=True,clean=True)
    if preserve_password and j.opt.get('pw'):
        opts.update(encryption=fitz.PDF_ENCRYPT_AES_256,user_pw=str(j.opt['pw']),owner_pw=str(j.opt['pw']))
    try:
        doc.save(tmp,**opts)
        with open_pdf(tmp,j.opt.get('pw','')) as check:
            if not len(check): raise ValueError('輸出 PDF 沒有頁面')
        os.replace(tmp,path); j.output=path;j.encoder='本機 PDF 引擎';j.validation='PDF 結構與頁數檢查通過'
    finally:
        if os.path.exists(tmp): os.remove(tmp)
def warn(j,message):
    if not hasattr(j,'notes'):j.notes=[]
    if message not in j.notes:j.notes.append(message)
def textpage(page,o):
    if o.get('ocr') or o.get('_force_ocr'):
        language=str(o.get('lang','chi_tra+eng'))
        if any(not re.fullmatch('[a-z_]+',v) or not os.path.isfile(os.path.join(tessdata(),v+'.traineddata')) for v in language.split('+')): raise ValueError('離線 OCR 語言資料不完整')
        return page.get_textpage_ocr(language=language,dpi=200,full=bool(o.get('_force_ocr')),tessdata=tessdata())
    return page.get_textpage()
def recognized_page(page,o,dpi=200):
    import numpy as np
    from PIL import Image
    language=str(o.get('lang','chi_tra+eng'))
    if any(not re.fullmatch('[a-z_]+',v) or not os.path.isfile(os.path.join(tessdata(),v+'.traineddata')) for v in language.split('+')):raise ValueError('離線 OCR 語言資料不完整')
    dpi=max(72,min(int(dpi),int(2600/max(page.rect.width,page.rect.height)*72)))
    original=page.get_pixmap(dpi=dpi,alpha=False,colorspace=fitz.csRGB)
    try:pixels=np.frombuffer(original.samples,dtype=np.uint8).reshape(original.height,original.width,3).copy()   # no PNG round trip
    except Exception:pixels=np.asarray(Image.open(io.BytesIO(original.tobytes('png'))).convert('RGB')).copy()
    dark=pixels.mean(axis=2)<150;h,w=dark.shape
    # Long straight rules confuse automatic OCR page segmentation. Remove them only
    # in the recognition image, then restore the original image in the output PDF.
    horizontal=np.zeros(h,dtype=np.int32);max_horizontal=horizontal.copy()
    for x in range(w):horizontal=np.where(dark[:,x],horizontal+1,0);max_horizontal=np.maximum(max_horizontal,horizontal)
    vertical=np.zeros(w,dtype=np.int32);max_vertical=vertical.copy()
    for y in range(h):vertical=np.where(dark[y],vertical+1,0);max_vertical=np.maximum(max_vertical,vertical)
    rows=np.where(max_horizontal>max(80,w*.25))[0];cols=np.where(max_vertical>max(60,h*.06))[0]
    pixels[rows,:,:]=255;pixels[:,cols,:]=255
    try:pix=fitz.Pixmap(fitz.csRGB,pixels.shape[1],pixels.shape[0],np.ascontiguousarray(pixels).tobytes(),False)
    except Exception:
        prepared=Image.fromarray(pixels);buffer=io.BytesIO();prepared.save(buffer,'PNG');pix=fitz.Pixmap(buffer.getvalue())
    pix.set_dpi(dpi,dpi)
    data=pix.pdfocr_tobytes(language=language,tessdata=tessdata())
    result=fitz.open('pdf',data)
    images=result[0].get_images(full=True)
    if images:result[0].replace_image(images[0][0],pixmap=original)
    return result
def page_data(page,o,kind='text'):
    if o.get('ocr') or o.get('_force_ocr'):
        with recognized_page(page,o) as recognized:return recognized[0].get_text(kind,sort=True)
    return page.get_text(kind,sort=True)
def page_text(page,o):return page_data(page,o)
def put_text(page,box,text,size=12,ink=(0,0,0),opacity=1,align=0):
    if len(str(text))>20000: raise ValueError('單個文字框內容過長')
    cjk=bool(re.search(r'[^\x00-\x7f]',str(text)))
    rc=page.insert_textbox(box,str(text),fontname='china-ts' if cjk else 'helv',fontsize=float(size),color=ink,align=align,fill_opacity=float(opacity))
    if rc<0: raise ValueError('文字超出框選範圍，請放大範圍或縮小字級')

@local
def preview(path,page_number=1,pw=''):
    with open_pdf(path,pw) as doc:
        n=max(0,min(len(doc)-1,int(page_number)-1)); page=doc[n]
        if page.rotation:page.remove_rotation()
        scale=min(1.7,1200/max(page.rect.width,page.rect.height));pix=page.get_pixmap(matrix=fitz.Matrix(scale,scale),alpha=False)
        widgets=[dict(name=w.field_name,type=w.field_type_string,value=w.field_value,rect=[w.rect.x0,w.rect.y0,w.rect.width,w.rect.height],choices=w.choice_values or []) for w in page.widgets() or ()]
        return dict(src='data:image/png;base64,'+base64.b64encode(pix.tobytes('png')).decode(),w=page.rect.width,h=page.rect.height,dur=0,page=n+1,count=len(doc),forms=widgets,name=os.path.basename(path))

def arrange(j,o):
    if j.tool=='p_merge':
        result=fitz.open()
        try:
            for i,path in enumerate(j.files):
                cancelled(j)
                with open_pdf(path,o.get('pw')) as doc: result.insert_pdf(doc,widgets=True)
                j.progress=(i+1)/len(j.files)*.85
            save(result,output(j),j)
        finally: result.close()
        return
    with open_pdf(j.files[0],o.get('pw')) as doc:
        selected=indices(o.get('pages'),len(doc),allow_repeat=True)
        if j.tool=='p_remove':
            if len(set(selected))==len(doc): raise ValueError('不能移除文件的所有頁面')
            doc.delete_pages(sorted(set(selected)))
        elif j.tool=='p_extract': doc.select(selected)
        elif o.get('mode')=='reverse':doc.select(list(reversed(range(len(doc)))))
        elif o.get('mode')=='blank':
            for n in sorted(set(selected),reverse=True):
                box=doc[n].rect;doc.new_page(pno=n+1,width=box.width,height=box.height)
        else:doc.select(selected)
        save(doc,output(j),j)

def compress(j,o):
    import pikepdf
    with open_pdf(j.files[0],o.get('pw')) as doc:
        level=o.get('level','balanced')
        if level!='lossless':
            max_side,quality=(1800,78) if level=='balanced' else (1100,55)
            seen=set()
            from PIL import Image
            for n,page in enumerate(doc):
                cancelled(j)
                for image in page.get_images(full=True):
                    xref=image[0]
                    if xref in seen or image[1]: continue
                    seen.add(xref)
                    try:
                        data=doc.extract_image(xref)
                        with Image.open(io.BytesIO(data['image'])) as im:
                            if im.width*im.height<40000:continue
                            im=im.convert('RGB');im.thumbnail((max_side,max_side));buf=io.BytesIO();im.save(buf,'JPEG',quality=quality,optimize=True)
                            if len(buf.getvalue())<len(data['image']):page.replace_image(xref,stream=buf.getvalue())
                    except (ValueError,RuntimeError,OSError):continue
                j.progress=(n+1)/len(doc)*.8
        path=output(j);save(doc,path,j)
    if os.path.getsize(path)>=os.path.getsize(j.files[0]):
        shutil.copyfile(j.files[0],path);j.validation='原檔較小，已保留原始位元內容';warn(j,'原始檔已較精簡，沒有強制降低品質以追求檔案大小')

def repair(j,o):
    report=dict(source=os.path.basename(j.files[0]),recovered=0,unreadable=[],method='')
    root=folder(j);dest=os.path.join(root,'修復後.pdf')
    try:
        import pikepdf
        with pikepdf.open(j.files[0],password=str(o.get('pw') or ''),attempt_recovery=True) as pdf:
            options=dict(object_stream_mode=pikepdf.ObjectStreamMode.generate)
            if o.get('pw'):options['encryption']=pikepdf.Encryption(owner=str(o['pw']),user=str(o['pw']),R=6)
            pdf.save(dest,**options)
        report['method']='qpdf 結構重建'
        with open_pdf(dest,o.get('pw')) as doc:
            for i,page in enumerate(doc):
                cancelled(j);page.get_pixmap(matrix=fitz.Matrix(.2,.2));report['recovered']+=1;j.progress=(i+1)/len(doc)*.9
    except InterruptedError:raise
    except Exception as first:
        report['first_attempt']=type(first).__name__
        with open_pdf(j.files[0],o.get('pw')) as doc:
            recovered=fitz.open()
            try:
                for i,page in enumerate(doc):
                    cancelled(j)
                    try:
                        page.get_pixmap(matrix=fitz.Matrix(.3,.3));recovered.insert_pdf(doc,from_page=i,to_page=i);report['recovered']+=1
                    except Exception:report['unreadable'].append(i+1)
                if not len(recovered):raise ValueError('無法恢復可讀取的頁面；原檔保持不變')
                options=dict(garbage=4,deflate=True)
                if o.get('pw'):options.update(encryption=fitz.PDF_ENCRYPT_AES_256,user_pw=str(o['pw']),owner_pw=str(o['pw']))
                recovered.save(dest,**options)
            finally:recovered.close()
        report['method']='MuPDF 逐頁復原'
    if report['unreadable']:j.warnings.append(f'{len(report["unreadable"])} 頁無法恢復，請查看報告')
    report['note']='可恢復的內容取決於原檔損壞程度，修復後請檢查頁數與畫面。'
    with open(os.path.join(root,'修復報告.json'),'w',encoding='utf-8') as f:json.dump(report,f,ensure_ascii=False,indent=2)
    j.encoder=report['method'];j.validation=f'逐頁驗證 {report["recovered"]} 頁'

def ocr_pdf(j,o):
    root=folder(j);result=fitz.open();texts=[]
    try:
        with open_pdf(j.files[0],o.get('pw')) as doc:
            for n,page in enumerate(doc):
                cancelled(j);j.msg=f'辨識第 {n+1} / {len(doc)} 頁'
                if len(page.get_text().strip())>30 and not o.get('force'):
                    result.insert_pdf(doc,from_page=n,to_page=n);text=page.get_text(sort=True)
                else:
                    dpi=min(int(o.get('dpi',200)),int(2400/max(page.rect.width,page.rect.height)*72))
                    with recognized_page(page,o,dpi) as recognized:
                        text=recognized[0].get_text(sort=True);result.insert_pdf(recognized)
                texts.append(text);j.progress=(n+1)/len(doc)*.9
        save(result,os.path.join(root,'可搜尋文件.pdf'),j)
        with open(os.path.join(root,'辨識文字.txt'),'w',encoding='utf-8-sig') as f:f.write('\n\f\n'.join(texts))
        j.output=root;j.encoder='離線 Tesseract OCR';j.validation='可搜尋文字層與 PDF 結構檢查完成';warn(j,'掃描頁面會重新產生影像與文字層，請檢查辨識結果與頁面細節')
    finally:result.close()

def clean_scan(im,o):
    from PIL import ImageOps,ImageEnhance
    im=ImageOps.exif_transpose(im).convert('RGB');im.thumbnail((3000,3000))
    if o.get('clean'):im=ImageOps.autocontrast(ImageOps.grayscale(im)).convert('RGB')
    if o.get('deskew'):
        import numpy as np
        sample=ImageOps.grayscale(im.copy());sample.thumbnail((650,650));best=(0,-1)
        for angle in (a*.5 for a in range(-10,11)):
            candidate=sample.rotate(angle,fillcolor=255)
            ink=np.asarray(candidate)<150
            score=float(np.var(ink.sum(axis=1)))
            if score>best[1]:best=(angle,score)
        if best[0]:im=im.rotate(best[0],expand=True,fillcolor='white')
    return im

def images_pdf(j,o):
    from PIL import Image,ImageOps
    doc=fitz.open()
    try:
        for n,path in enumerate(j.files):
            cancelled(j)
            with Image.open(path) as source:
                im=clean_scan(source,o) if j.tool=='p_scan' else ImageOps.exif_transpose(source).convert('RGB')
                im.thumbnail((4000,4000));buf=io.BytesIO();im.save(buf,'JPEG',quality=90)
                paper=o.get('paper','a4');w,h=(612,792) if paper=='letter' else (595,842) if paper=='a4' else (im.width*.75,im.height*.75)
                if o.get('landscape'):w,h=h,w
                margin=0 if paper=='original' else max(0,min(float(o.get('margin',10)),50))*72/25.4
                page=doc.new_page(width=w,height=h);page.insert_image(fitz.Rect(margin,margin,w-margin,h-margin),stream=buf.getvalue())
                if o.get('ocr'):
                    with recognized_page(page,o,150) as replacement:
                        doc.delete_page(len(doc)-1);doc.insert_pdf(replacement)
            j.progress=(n+1)/len(j.files)*.9
        save(doc,output(j),j)
    finally:doc.close()

def convert_office(eng,j,o):
    extensions={'p_word_pdf':{'.doc','.docx','.odt','.rtf'},'p_ppt_pdf':{'.ppt','.pptx','.odp'},'p_excel_pdf':{'.xls','.xlsx','.ods','.csv'},'p_html_pdf':{'.html','.htm','.xhtml'}}
    allowed=extensions.get(j.tool)
    if allowed and os.path.splitext(j.files[0])[1].lower() not in allowed:raise ValueError('請加入此工具對應的文件格式')
    if not E.SOFFICE:raise ValueError('離線完整版本需要隨附 LibreOffice 元件')
    if j.tool=='p_html_pdf':
        text=open(j.files[0],encoding='utf-8-sig',errors='replace').read()
        if re.search(r'<script\b|<iframe\b|<object\b|<embed\b|@import\b|(?:https?|ftp):|(?:src|href)\s*=\s*[\"\']?(?://|\\\\)|url\(\s*[\"\']?(?://|\\\\)',text,re.I):
            raise ValueError('HTML 含網路資源，請先將圖片與樣式儲存至本機再轉換')
    previous=j.target;j.target='pdf'
    try:eng._doc(j)
    finally:j.target=previous

def pdf_images(j,o):
    root=folder(j);seen=set();count=0
    with open_pdf(j.files[0],o.get('pw')) as doc:
        selected=indices(o.get('pages'),len(doc))
        for i,n in enumerate(selected):
            cancelled(j);page=doc[n]
            if o.get('mode')=='embedded':
                for image in page.get_images(full=True):
                    if image[0] in seen:continue
                    seen.add(image[0]);data=doc.extract_image(image[0]);count+=1
                    with open(os.path.join(root,f'圖片_{count:04d}.{data["ext"]}'),'wb') as f:f.write(data['image'])
            else:
                dpi=min(int(o.get('dpi',150)),int(4000/max(page.rect.width,page.rect.height)*72))
                pix=page.get_pixmap(dpi=max(72,dpi),alpha=False);ext=o.get('fmt','jpg');pix.save(os.path.join(root,f'第{n+1:04d}頁.{ext}'));count+=1
            j.progress=(i+1)/len(selected)*.95
    if not count:raise ValueError('文件沒有可提取的圖片')
    j.encoder='本機頁面渲染';j.validation=f'輸出 {count} 張圖片'

def to_word(j,o):
    from docx import Document
    from docx.shared import Inches
    document=Document();any_text=False
    if j.tool.endswith('_ocr'):o=dict(o,ocr=True,_force_ocr=True)
    with open_pdf(j.files[0],o.get('pw')) as pdf:
        for n,page in enumerate(pdf):
            cancelled(j);text=page_text(page,o);any_text=any_text or bool(text.strip())
            for paragraph in re.split(r'\n\s*\n',text.strip()):
                if paragraph:document.add_paragraph(paragraph)
            if not o.get('ocr'):
                for image in page.get_images(full=True):
                    try:document.add_picture(io.BytesIO(pdf.extract_image(image[0])['image']),width=Inches(5.5))
                    except Exception:continue
                for table in page.find_tables().tables:
                    rows=table.extract()
                    if rows and max(map(len,rows)):
                        grid=document.add_table(rows=0,cols=max(map(len,rows)));grid.style='Table Grid'
                        for values in rows:
                            cells=grid.add_row().cells
                            for i,value in enumerate(values):cells[i].text=str(value or '')
            if n<len(pdf)-1:document.add_page_break()
            j.progress=(n+1)/len(pdf)*.9
    if not any_text:raise ValueError('沒有可編輯文字，請啟用 OCR 或使用 PDF 轉 Word（OCR）')
    path=output(j,'docx');document.save(path);j.output=path;j.encoder='本機 Word 重建';warn(j,'PDF 的文字、圖片與表格已重建為 Word；複雜排版不保證與原頁面相同')

def to_ppt(j,o):
    from pptx import Presentation
    from pptx.util import Pt
    prs=Presentation()
    with open_pdf(j.files[0],o.get('pw')) as pdf:
        prs.slide_width=int(pdf[0].rect.width*12700);prs.slide_height=int(pdf[0].rect.height*12700)
        for n,page in enumerate(pdf):
            cancelled(j);slide=prs.slides.add_slide(prs.slide_layouts[6]);sx=prs.slide_width/page.rect.width;sy=prs.slide_height/page.rect.height
            if o.get('mode')=='visual':
                pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False);slide.shapes.add_picture(io.BytesIO(pix.tobytes('png')),0,0,width=prs.slide_width,height=prs.slide_height)
            else:
                data=page_data(page,o,'dict')
                for block in data['blocks']:
                    if block['type']!=0:continue
                    x,y,x1,y1=block['bbox'];box=slide.shapes.add_textbox(int(x*sx),int(y*sy),max(1,int((x1-x)*sx)),max(1,int((y1-y)*sy)))
                    frame=box.text_frame;frame.word_wrap=True;frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
                    for i,line in enumerate(block['lines']):
                        paragraph=frame.paragraphs[0] if i==0 else frame.add_paragraph()
                        for span in line['spans']:
                            run=paragraph.add_run();run.text=span['text'];run.font.size=Pt(max(6,min(72,span['size'])))
                for image in page.get_image_info(xrefs=True):
                    if not image.get('xref'):continue
                    try:
                        x,y,x1,y1=image['bbox'];slide.shapes.add_picture(io.BytesIO(pdf.extract_image(image['xref'])['image']),int(x*sx),int(y*sy),width=max(1,int((x1-x)*sx)),height=max(1,int((y1-y)*sy)))
                    except Exception:continue
            j.progress=(n+1)/len(pdf)*.9
    path=output(j,'pptx');prs.save(path);j.output=path;j.encoder='本機 PowerPoint 重建';warn(j,'視覺模式是整頁圖片；可編輯模式的文字與圖片可調整，向量繪圖與複雜版面可能有差異')

def excel_cell(value):
    value=str(value or '')
    return "'"+value if value.startswith(('=','+','-','@')) else value
def to_excel(j,o):
    from openpyxl import Workbook
    book=Workbook();book.remove(book.active);found=0
    if j.tool.endswith('_ocr'):o=dict(o,ocr=True,_force_ocr=True)
    with open_pdf(j.files[0],o.get('pw')) as pdf:
        for n,page in enumerate(pdf):
            cancelled(j)
            tables=[] if o.get('ocr') else [t.extract() for t in page.find_tables().tables]
            if tables:
                for index,rows in enumerate(tables,1):
                    sheet=book.create_sheet(f'頁{n+1}_表{index}');found+=1
                    for row in rows:sheet.append([excel_cell(v) for v in row])
            else:
                words=page_data(page,o,'words');lines={}
                for x0,y0,x1,y1,text,b,l,w in words:lines.setdefault((b,l),[]).append((x0,x1,text))
                sheet=book.create_sheet(f'頁{n+1}_文字')
                for line in lines.values():
                    cells=[];last=None
                    for x0,x1,text in sorted(line):
                        if last is None or x0-last>max(12,(y1-y0)*1.3):cells.append(text)
                        else:cells[-1]+=' '+text
                        last=x1
                    sheet.append([excel_cell(v) for v in cells])
            j.progress=(n+1)/len(pdf)*.9
    if not any(sheet.max_row and any(c.value for row in sheet for c in row) for sheet in book):raise ValueError('沒有可擷取的文字／表格，請啟用 OCR')
    path=output(j,'xlsx');book.save(path);j.output=path;j.encoder='本機表格重建';j.validation=f'{found} 個文字表格；其餘頁面使用欄位位置重建';warn(j,'無框線及掃描表格的欄位判斷可能有誤，請檢查資料；原文的公式樣式字串以文字儲存')

def numbering(j,o):
    with open_pdf(j.files[0],o.get('pw')) as doc:
        chosen=indices(o.get('pages'),len(doc));size=max(6,min(72,float(o.get('size',11))))
        for i,n in enumerate(chosen):
            cancelled(j);page=doc[n]
            if page.rotation:page.remove_rotation()
            pos=o.get('pos','bc');y=20 if pos.startswith('t') else page.rect.height-size*2-14
            number=int(o.get('start',1))+i;label=f'{number} / {int(o.get("start",1))+len(chosen)-1}' if o.get('format')=='total' else str(number)
            put_text(page,fitz.Rect(24,y,page.rect.width-24,y+size*2),label,size,color(o.get('color','#333333')),align={'l':0,'c':1,'r':2}.get(pos[-1],1))
            j.progress=(i+1)/len(chosen)*.9
        save(doc,output(j),j)

def watermark(j,o):
    with open_pdf(j.files[0],o.get('pw')) as doc:
        chosen=indices(o.get('pages'),len(doc));opacity=max(.01,min(1,float(o.get('opacity',25))/100));size=max(5,min(90,float(o.get('size',35))))/100
        for i,n in enumerate(chosen):
            cancelled(j);page=doc[n]
            if page.rotation:page.remove_rotation()
            w=page.rect.width*size;h=max(35,w*.22);spots=[((page.rect.width-w)/2,(page.rect.height-h)/2)]
            if o.get('tile'):spots=[(x,y) for y in range(40,int(page.rect.height)-int(h),int(h+70)) for x in range(20,int(page.rect.width)-int(w),int(w+30))]
            for x,y in spots:
                box=fitz.Rect(x,y,x+w,y+h)
                if o.get('kind')=='image':
                    from PIL import Image
                    with Image.open(o.get('image','')) as image:
                        im=image.convert('RGBA');im.thumbnail((1500,1500));alpha=im.getchannel('A').point(lambda a:int(a*opacity));im.putalpha(alpha);buf=io.BytesIO();im.save(buf,'PNG');page.insert_image(box,stream=buf.getvalue())
                else:
                    text=str(o.get('text') or '')
                    if not text:raise ValueError('請輸入浮水印文字')
                    font=fitz.Font('china-ts' if re.search(r'[^\x00-\x7f]',text) else 'helv');fs=min(h*.65,w/max(font.text_length(text,fontsize=1),1))
                    page.insert_text((x,y+h*.75),text,fontname='china-ts' if re.search(r'[^\x00-\x7f]',text) else 'helv',fontsize=fs,color=color(o.get('color','#808080')),fill_opacity=opacity,
                                     morph=(fitz.Point(x+w/2,y+h/2),fitz.Matrix(25)) if o.get('tile') else None)
            j.progress=(i+1)/len(chosen)*.9
        save(doc,output(j),j)

def crop(j,o):
    with open_pdf(j.files[0],o.get('pw')) as doc:
        for n in indices(o.get('pages'),len(doc)):
            cancelled(j);page=doc[n]
            if page.rotation:page.remove_rotation()
            if o.get('mode')=='rect':box=rect(o.get('rect'),page)
            else:
                m=float(o.get('margin',10))*72/25.4;box=fitz.Rect(m,m,page.rect.width-m,page.rect.height-m)
                if box.is_empty:raise ValueError('邊距過大')
            # Coordinates in the UI are relative to the visible crop box.
            offset=page.cropbox_position
            page.set_cropbox(box+(offset.x,offset.y,offset.x,offset.y))
        save(doc,output(j),j);warn(j,'裁切只改變可見範圍；若要永久刪除敏感內容，請使用永久遮蔽')

def edit(j,o):
    actions=safe_actions(o.get('actions'))
    if not actions:raise ValueError('請在 PDF 工作台加入編輯項目')
    with open_pdf(j.files[0],o.get('pw')) as doc:
        for action in actions:
            cancelled(j);page=selected_page(doc,action);kind=action.get('kind');box=rect(action.get('rect'),page,action.get('normalized',False));ink=color(action.get('color','#000000'))
            if kind in ('text','date','name'):
                put_text(page,box,action.get('text') or time.strftime('%Y-%m-%d'),action.get('size',12),ink)
            elif kind in ('image','signature'):
                if action.get('data'):
                    data=str(action['data']).split(',',1)[-1]
                    if len(data)>3*1024*1024:raise ValueError('簽名圖片過大')
                    stream=base64.b64decode(data,validate=True)
                    from PIL import Image
                    with Image.open(io.BytesIO(stream)) as image:image.verify()
                    page.insert_image(box,stream=stream)
                else:page.insert_image(box,filename=str(action.get('file') or ''))
            elif kind=='rectangle':page.draw_rect(box,color=ink,width=max(.1,min(20,float(action.get('width',1.5)))))
            elif kind=='highlight':page.draw_rect(box,fill=ink,fill_opacity=.2,color=None)
            elif kind=='draw':
                points=action.get('points') or []
                if not 2<=len(points)<=5000:raise ValueError('手繪筆跡資料無效')
                coords=[fitz.Point(float(x)*page.rect.width,float(y)*page.rect.height) for x,y in points]
                if any(not page.rect.contains(p) for p in coords):raise ValueError('手繪超出頁面')
                page.draw_polyline(coords,color=ink,width=max(.1,min(20,float(action.get('width',1.5)))))
            else:raise ValueError('不支援的編輯類型')
        save(doc,output(j),j)

def forms(j,o):
    actions=safe_actions(o.get('actions'));mode=o.get('mode','fill')
    with open_pdf(j.files[0],o.get('pw')) as doc:
        if mode=='flatten':doc.bake(annots=True,widgets=True)
        elif mode=='detect':
            count=0
            for n,page in enumerate(doc):
                cancelled(j)
                for drawing in page.get_drawings():
                    for item in drawing['items']:
                        candidate=None
                        if item[0]=='re' and 45<item[1].width<page.rect.width*.9 and 8<item[1].height<70:candidate=item[1]
                        elif item[0]=='l' and abs(item[1].y-item[2].y)<1 and abs(item[1].x-item[2].x)>50:candidate=fitz.Rect(min(item[1].x,item[2].x),item[1].y-18,max(item[1].x,item[2].x),item[1].y-1)
                        if candidate and page.rect.contains(candidate):
                            count+=1;widget=fitz.Widget();widget.field_name=f'候選欄位_{n+1}_{count}';widget.field_type=fitz.PDF_WIDGET_TYPE_TEXT;widget.rect=candidate;widget.text_font='Helv';widget.text_fontsize=10;page.add_widget(widget)
                            if count>=500:break
            if not count:raise ValueError('沒有找到底線或方框，請在工作台手動建立欄位')
            warn(j,f'已建立 {count} 個候選欄位，請人工檢查欄位位置與名稱')
        elif mode=='create':
            if not actions:raise ValueError('請在工作台加入表單欄位')
            existing={w.field_name for p in doc for w in p.widgets() or ()}
            for action in actions:
                cancelled(j);page=selected_page(doc,action);name=str(action.get('name') or '').strip()[:120]
                if not name or name in existing:raise ValueError('表單欄位名稱不可空白或重複')
                existing.add(name);widget=fitz.Widget();widget.field_name=name;widget.rect=rect(action.get('rect'),page,action.get('normalized',False))
                kind=action.get('field_type','text');widget.field_type={'text':fitz.PDF_WIDGET_TYPE_TEXT,'checkbox':fitz.PDF_WIDGET_TYPE_CHECKBOX,'choice':fitz.PDF_WIDGET_TYPE_COMBOBOX,'list':fitz.PDF_WIDGET_TYPE_LISTBOX,'radio':fitz.PDF_WIDGET_TYPE_RADIOBUTTON}.get(kind,fitz.PDF_WIDGET_TYPE_TEXT)
                widget.text_font='Helv';widget.text_fontsize=10;widget.border_width=1;widget.border_color=(.5,.5,.5);widget.fill_color=(1,1,1)
                if kind in ('choice','list'):widget.choice_values=[str(v) for v in action.get('choices',[])][:100]
                widget.field_value=action.get('value','Off' if kind in ('checkbox','radio') else '')
                page.add_widget(widget)
        else:
            values={str(a.get('name')):a.get('value','') for a in actions};count=0
            if not values:raise ValueError('請在工作台填寫表單欄位')
            for page in doc:
                for widget in page.widgets() or ():
                    if widget.field_name in values:
                        value=values[widget.field_name]
                        if widget.field_type in (fitz.PDF_WIDGET_TYPE_CHECKBOX,fitz.PDF_WIDGET_TYPE_RADIOBUTTON):value=widget.on_state() if value in (True,1,'true','Yes','On') else 'Off'
                        widget.field_value=value;widget.update();count+=1
            if not count:raise ValueError('找不到指定的表單欄位')
        if o.get('flatten'):doc.bake(annots=True,widgets=True)
        save(doc,output(j),j)

def redact(j,o):
    if not o.get('confirmed'):raise ValueError('請先確認遮蔽範圍，再勾選永久移除內容')
    actions=safe_actions(o.get('actions'));text=str(o.get('text') or '');count=0
    with open_pdf(j.files[0],o.get('pw')) as doc:
        chosen=indices(o.get('pages'),len(doc))
        for n in chosen:
            cancelled(j);page=doc[n]
            if page.rotation:page.remove_rotation()
            boxes=[]
            if text:boxes+=page.search_for(text)
            if o.get('rect') and n==int(o.get('preview_page',1))-1:boxes.append(rect(o['rect'],page))
            boxes += [rect(a['rect'],page,a.get('normalized',False)) for a in actions if int(a.get('page',1))-1==n]
            for box in boxes:page.add_redact_annot(box,fill=(0,0,0));count+=1
            if boxes:
                page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS,graphics=fitz.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED,text=fitz.PDF_REDACT_TEXT_REMOVE)
                # Remove overlays, hidden notes and form values from a redacted page.
                for annotation in list(page.annots() or ()):page.delete_annot(annotation)
                for widget in list(page.widgets() or ()):page.delete_widget(widget)
        if not count:raise ValueError('沒有找到要遮蔽的文字或範圍')
        doc.set_metadata({})
        if doc.xref_xml_metadata():doc.del_xml_metadata()
        for name in list(doc.embfile_names()):doc.embfile_del(name)
        doc.set_toc([])
        save(doc,output(j),j)
        with open_pdf(j.output,o.get('pw')) as check:
            if text and any(text in check[n].get_text() for n in chosen):raise ValueError('文字移除驗證失敗，請擴大遮蔽範圍')
        j.validation=f'{count} 個範圍永久移除；清除註記、欄位、附件與中繼資料';warn(j,'請檢查其他頁面、圖片內的文字及文件副本；未選取範圍不會自動遮蔽')

def compare(j,o):
    if len(j.files)!=2:raise ValueError('比較 PDF 請選擇兩份文件')
    from PIL import Image,ImageChops,ImageStat
    root=folder(j);records=[];cards=[]
    with open_pdf(j.files[0],o.get('pw')) as first,open_pdf(j.files[1],o.get('pw')) as second:
        for n in range(max(len(first),len(second))):
            cancelled(j);texts=[doc[n].get_text(sort=True) if n<len(doc) else '' for doc in (first,second)]
            diff='\n'.join(difflib.unified_diff(texts[0].splitlines(),texts[1].splitlines(),fromfile='原版本',tofile='新版本',lineterm=''));changed=bool(diff);visual=None
            if o.get('visual',True):
                images=[]
                for index,doc in enumerate((first,second)):
                    if n<len(doc):
                        page=doc[n];pix=page.get_pixmap(matrix=fitz.Matrix(min(1,900/max(page.rect.width,page.rect.height)),min(1,900/max(page.rect.width,page.rect.height))),alpha=False)
                        image=Image.open(io.BytesIO(pix.tobytes('png'))).convert('RGB')
                    else:image=Image.new('RGB',(600,800),'white')
                    images.append(image);image.save(os.path.join(root,f'p{n+1}_{index}.png'))
                w=max(im.width for im in images);h=max(im.height for im in images);canvases=[]
                for image in images:
                    canvas=Image.new('RGB',(w,h),'white');canvas.paste(image,(0,0));canvases.append(canvas)
                delta=ImageChops.difference(*canvases);visual=sum(ImageStat.Stat(delta).mean)/3;delta.save(os.path.join(root,f'p{n+1}_diff.png'));changed=changed or visual>.1
            record=dict(page=n+1,changed=changed,text_changed=bool(diff),visual_difference=visual);records.append(record)
            cards.append(f'<section><h2>第 {n+1} 頁：{"有差異" if changed else "一致"}</h2>'+ (f'<div class="grid"><img src="p{n+1}_0.png"><img src="p{n+1}_1.png"><img src="p{n+1}_diff.png"></div>' if visual is not None else '')+f'<pre>{html.escape(diff or "文字一致")}</pre></section>')
            j.progress=(n+1)/max(len(first),len(second))*.95
    with open(os.path.join(root,'比較報告.html'),'w',encoding='utf-8') as f:f.write('<!doctype html><meta charset="utf-8"><title>PDF 比較報告</title><style>body{font-family:Segoe UI,sans-serif;margin:24px}.grid{display:grid;grid-template-columns:repeat(3,1fr);gap:12px}img{width:100%;border:1px solid #ddd}pre{white-space:pre-wrap}section{border-bottom:1px solid #ddd;padding:12px}</style><h1>PDF 並排比較</h1>'+''.join(cards))
    with open(os.path.join(root,'比較報告.json'),'w',encoding='utf-8') as f:json.dump(records,f,ensure_ascii=False,indent=2)
    j.encoder='本機文字與像素比較';j.validation=f'{sum(r["changed"] for r in records)} 頁有差異；字型渲染微差也可能被標記'

def markdown(j,o):
    root=folder(j);parts=[]
    with open_pdf(j.files[0],o.get('pw')) as doc:
        for n,page in enumerate(doc):
            cancelled(j);parts.append(f'\n## 第 {n+1} 頁\n')
            data=page_data(page,o,'dict')
            sizes=[s['size'] for b in data['blocks'] if b['type']==0 for l in b['lines'] for s in l['spans']];median=sorted(sizes)[len(sizes)//2] if sizes else 12
            for block in data['blocks']:
                if block['type']!=0:continue
                for line in block['lines']:
                    text=''.join(s['text'] for s in line['spans']).strip()
                    if not text:continue
                    size=max((s['size'] for s in line['spans']),default=median)
                    parts.append(('### ' if size>median*1.3 and len(text)<150 else '')+text+'\n')
                parts.append('\n')
            if not o.get('ocr'):
                for table in page.find_tables().tables:
                    rows=table.extract()
                    if not rows:continue
                    cols=max(map(len,rows));parts.append('\n'+'| '+' | '.join(str(v or '').replace('|','\\|').replace('\n',' ') for v in rows[0])+ ' |\n| '+' | '.join(['---']*cols)+' |\n')
                    for row in rows[1:]:parts.append('| '+' | '.join(str(v or '').replace('|','\\|').replace('\n',' ') for v in row)+ ' |\n')
            for index,image in enumerate(page.get_images(full=True)):
                info=doc.extract_image(image[0]);name=f'page{n+1}_image{index+1}.{info["ext"]}'
                with open(os.path.join(root,name),'wb') as f:f.write(info['image'])
                parts.append(f'\n![第 {n+1} 頁圖片]({name})\n')
            for link in page.get_links():
                uri=link.get('uri','')
                if uri.startswith(('https://','http://')):parts.append(f'\n[連結]({uri.replace(")","%29")})\n')
            j.progress=(n+1)/len(doc)*.9
    with open(os.path.join(root,'文件.md'),'w',encoding='utf-8') as f:f.write(''.join(parts))
    j.encoder='本機結構擷取';warn(j,'PDF 無明確語意時，標題、表格與閱讀順序以版面推測，請檢查')

def content(j,o):
    root=folder(j);info=[];texts=[]
    with open_pdf(j.files[0],o.get('pw')) as doc:
        for n,page in enumerate(doc):
            cancelled(j);texts.append(page.get_text(sort=True));info.append(dict(page=n+1,width=page.rect.width,height=page.rect.height,rotation=page.rotation))
            for index,image in enumerate(page.get_images(full=True)):
                data=doc.extract_image(image[0])
                with open(os.path.join(root,f'頁{n+1}_圖片{index+1}.{data["ext"]}'),'wb') as f:f.write(data['image'])
        attachment_names=[]
        for index,name in enumerate(doc.embfile_names()):
            safe=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',os.path.basename(name))[:100] or '附件';safe=f'附件_{index+1}_{safe}'
            with open(os.path.join(root,safe),'wb') as f:f.write(doc.embfile_get(name))
            attachment_names.append(safe)
        report=dict(pages=info,metadata=doc.metadata,bookmarks=doc.get_toc(),attachments=attachment_names)
        with open(os.path.join(root,'文件資訊.json'),'w',encoding='utf-8') as f:json.dump(report,f,ensure_ascii=False,indent=2)
    with open(os.path.join(root,'文字.txt'),'w',encoding='utf-8-sig') as f:f.write('\n\f\n'.join(texts))
    j.encoder='本機內容擷取'

def certificate(j,o):
    from pyhanko.sign import signers
    from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
    pfx=str(o.get('pfx') or '')
    if not os.path.isfile(pfx):raise ValueError('請選擇 PFX／P12 憑證')
    with open(pfx,'rb') as f:key=signers.SimpleSigner.load_pkcs12_data(f.read(),other_certs=[],passphrase=str(o.get('pfx_pw') or '').encode() or None)
    path=output(j)
    with open(j.files[0],'rb') as source,open(path,'wb') as dest:
        signer=signers.PdfSigner(signers.PdfSignatureMetadata(field_name=str(o.get('name') or 'Signature1'),reason=str(o.get('reason') or '')),signer=key)
        signer.sign_pdf(IncrementalPdfFileWriter(source),output=dest)
    j.output=path;j.encoder='本機憑證簽章';warn(j,'沒有線上時間戳或撤銷狀態查詢；信任程度取決於你的憑證與驗證者設定，簽署後再編輯會影響簽章')

def run(eng,j,o=None):
    if j.tool in ('p_summary','p_translate','p_pdfa','p_workflow'):
        import pdf_local_engines
        if j.tool=='p_workflow':return pdf_local_engines.workflow(eng,j,dict(o or j.opt))
        return pdf_local_engines.pdfa(eng,j,dict(o or j.opt)) if j.tool=='p_pdfa' else pdf_local_engines.intelligence(eng,j,dict(o or j.opt))
    return _run_local(eng,j,o)

@local
def _run_local(eng,j,o=None):
    o=dict(o or j.opt);cancelled(j)
    if j.tool in ('p_merge','p_remove','p_extract','p_organize'):arrange(j,o)
    elif j.tool=='p_compress':compress(j,o)
    elif j.tool=='p_repair':repair(j,o)
    elif j.tool=='p_ocr':ocr_pdf(j,o)
    elif j.tool in ('p_images_pdf','p_scan'):images_pdf(j,o)
    elif j.tool in ('p_office_pdf','p_word_pdf','p_ppt_pdf','p_excel_pdf','p_html_pdf'):convert_office(eng,j,o)
    elif j.tool=='p_pdf_images':pdf_images(j,o)
    elif j.tool in ('p_pdf_word','p_pdf_word_ocr'):to_word(j,o)
    elif j.tool=='p_pdf_ppt':to_ppt(j,o)
    elif j.tool in ('p_pdf_excel','p_pdf_excel_ocr'):to_excel(j,o)
    elif j.tool=='p_numbers':numbering(j,o)
    elif j.tool=='p_watermark':watermark(j,o)
    elif j.tool=='p_crop':crop(j,o)
    elif j.tool in ('p_edit','p_sign'):edit(j,o)
    elif j.tool=='p_forms':forms(j,o)
    elif j.tool=='p_redact':redact(j,o)
    elif j.tool=='p_compare':compare(j,o)
    elif j.tool=='p_markdown':markdown(j,o)
    elif j.tool=='p_content':content(j,o)
    elif j.tool=='p_certificate':certificate(j,o)
    elif j.tool=='p_copy':
        path=output(j);shutil.copyfile(j.files[0],path);j.output=path;j.encoder='完整副本'
        with open(j.files[0],'rb') as original,open(path,'rb') as duplicate:
            if hashlib.file_digest(original,'sha256').digest()!=hashlib.file_digest(duplicate,'sha256').digest():raise ValueError('複製驗證失敗')
    elif j.tool=='p_flatten':
        with open_pdf(j.files[0],o.get('pw')) as doc:doc.bake(annots=True,widgets=True);save(doc,output(j),j)
    elif j.tool=='p_pdfa':
        import pdf_local_engines;pdf_local_engines.pdfa(eng,j,o)
    elif j.tool in ('p_summary','p_translate'):
        import pdf_local_engines;pdf_local_engines.intelligence(eng,j,o)
    else:raise ValueError('未知 PDF 工具')
    cancelled(j)
