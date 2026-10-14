import os,sys,io,json,tempfile,unittest,hashlib,threading,http.server,urllib.parse,faulthandler

from pathlib import Path
from unittest.mock import patch
ROOT=Path(__file__).resolve().parent
sys.path[:0]=[str(ROOT/'src'),'E:/四維八德/_build/vendor']
(ROOT/'scratch').mkdir(exist_ok=True);tempfile.tempdir=str(ROOT/'scratch')
import engine as E
E.base_dir=lambda:str(ROOT/'release');E._DATA=str(ROOT/'scratch/data');os.makedirs(E._DATA,exist_ok=True)
E.FFMPEG=E.find_tool('ffmpeg');E.SEVENZIP=E.find_tool('7z');E.SOFFICE=E.find_soffice();E.PANDOC=E.find_tool('pandoc')
import pdf_tools as P,pdf_schema as PS,tools as T,file_actions as FA,capture as C,workflows as W
import pymupdf as fitz
from PIL import Image

class PDFSuite(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.tmp=tempfile.TemporaryDirectory(prefix='pdf_checks_');cls.root=Path(cls.tmp.name);cls.source=cls.root/'測試原稿.pdf'
  cls.image=cls.root/'照片.png';Image.new('RGB',(600,800),'white').save(cls.image)
  doc=fitz.open();p=doc.new_page(width=400,height=600);p.insert_text((40,70),'CONFIDENTIAL SECRET 12345',fontsize=14);p.insert_text((40,120),'Public information for offline testing.',fontsize=12)
  widget=fitz.Widget();widget.field_name='name';widget.field_type=fitz.PDF_WIDGET_TYPE_TEXT;widget.rect=fitz.Rect(40,200,220,230);widget.field_value='Original';widget.text_font='Helv';p.add_widget(widget)
  p=doc.new_page(width=400,height=600)
  for y in (60,90,120):p.draw_line((40,y),(300,y))
  for x in (40,170,300):p.draw_line((x,60),(x,120))
  for point,text in [((50,80),'Item'),((180,80),'Price'),((50,110),'Apple'),((180,110),'20')]:p.insert_text(point,text,fontsize=12)
  p.insert_text((40,160),'=1+1',fontsize=12)
  p.insert_image(fitz.Rect(40,220,140,350),filename=str(cls.image))
  p=doc.new_page(width=400,height=600);p.insert_text((40,90),'Third page content.',fontsize=14)
  doc.set_metadata({'title':'SECRET metadata'});doc.embfile_add('note.txt',b'SECRET attachment');doc.save(cls.source);doc.close()
  with fitz.open(cls.source) as original:
   pix=original[1].get_pixmap(dpi=180,alpha=False);scanned=fitz.open();page=scanned.new_page(width=400,height=600);page.insert_image(page.rect,stream=pix.tobytes('png'));cls.scan=cls.root/'掃描稿.pdf';scanned.save(cls.scan);scanned.close()
  cls.source_hash=hashlib.sha256(cls.source.read_bytes()).hexdigest();cls.eng=E.Engine(1,light=1)
 @classmethod
 def tearDownClass(cls):cls.eng.shutdown();cls.tmp.cleanup()
 def execute(self,id,files=None,**options):
  target=self.root/('output_'+id);target.mkdir(exist_ok=True)
  defaults={f['k']:f['default'] for f in T.TOOL[id]['fields']};opt=dict(preset='balanced',out_dir=str(target),resource_mode='quiet',**defaults);opt.update(options)
  j=E.Job([str(p) for p in (files or [self.source])],id,opt,tool=id);j.state='running';T.run(self.eng,j);self.eng.validate(j)
  self.assertTrue(Path(j.output).exists());self.assertEqual(self.source_hash,hashlib.sha256(self.source.read_bytes()).hexdigest());return j
 def test_01_inventory(self):
  self.assertTrue(all(t['id'] in T.RUN or t['id']=='workflow' for t in PS.CATALOG));self.assertTrue(P.capabilities()['ocr']);self.assertTrue(P.capabilities()['ai']);self.assertTrue(P.capabilities()['pdfa'])
 def test_02_merge(self):
  j=self.execute('p_merge',[self.source,self.source]);self.assertEqual(len(fitz.open(j.output)),6)
 def test_03_remove(self):
  j=self.execute('p_remove',pages='2');self.assertEqual(len(fitz.open(j.output)),2)
 def test_04_extract_reorder(self):
  j=self.execute('p_extract',pages='3,1,1');doc=fitz.open(j.output);self.assertEqual(len(doc),3);self.assertIn('Third',doc[0].get_text());doc.close()
 def test_05_reorder_reverse(self):
  j=self.execute('p_organize',mode='reverse');doc=fitz.open(j.output);self.assertIn('Third',doc[0].get_text());doc.close()
 def test_06_blank_page(self):
  j=self.execute('p_organize',mode='blank',pages='1');doc=fitz.open(j.output);self.assertEqual(len(doc),4);self.assertEqual(doc[1].get_text(),'');doc.close()
 def test_07_remove_all_rejected(self):
  with self.assertRaises(ValueError):self.execute('p_remove',pages='1-')
 def test_08_compress(self):
  j=self.execute('p_compress',level='small');self.assertLessEqual(Path(j.output).stat().st_size,self.source.stat().st_size)
 def test_09_repair(self):
  damaged=self.root/'damaged.pdf';data=self.source.read_bytes();damaged.write_bytes(data.rsplit(b'startxref',1)[0]+b'%%EOF')
  j=self.execute('p_repair',[damaged]);self.assertEqual(json.loads((Path(j.output)/'修復報告.json').read_text(encoding='utf-8'))['recovered'],3)
 def test_10_ocr_searchable(self):
  j=self.execute('p_ocr',[self.scan],lang='eng');doc=fitz.open(str(Path(j.output)/'可搜尋文件.pdf'));self.assertIn('Apple',doc[0].get_text());doc.close()
 def test_11_images(self):
  j=self.execute('p_images_pdf',[self.image],margin=5);self.assertEqual(len(fitz.open(j.output)),1)
 def test_12_scan(self):
  j=self.execute('p_scan',[self.image],ocr=False);self.assertEqual(len(fitz.open(j.output)),1)
 def test_13_render(self):
  j=self.execute('p_pdf_images',fmt='png',pages='1,3');self.assertEqual(len(list(Path(j.output).glob('*.png'))),2)
 def test_14_extract_images(self):
  j=self.execute('p_pdf_images',mode='embedded');self.assertTrue(list(Path(j.output).glob('*.*')))
 def test_15_word(self):
  from docx import Document
  j=self.execute('p_pdf_word');self.assertIn('Public information','\n'.join(p.text for p in Document(j.output).paragraphs))
 def test_16_word_ocr(self):
  from docx import Document
  j=self.execute('p_pdf_word_ocr',[self.scan],lang='eng');self.assertIn('Apple','\n'.join(p.text for p in Document(j.output).paragraphs))
 def test_17_ppt(self):
  from pptx import Presentation
  j=self.execute('p_pdf_ppt');prs=Presentation(j.output);self.assertEqual(len(prs.slides),3);self.assertTrue(any(shape.has_text_frame and 'Public' in shape.text for shape in prs.slides[0].shapes))
 def test_18_excel(self):
  from openpyxl import load_workbook
  j=self.execute('p_pdf_excel');book=load_workbook(j.output);values=[str(c.value) for sheet in book for row in sheet for c in row];self.assertIn('Apple',values);self.assertFalse(any(c.data_type=='f' for sheet in book for row in sheet for c in row))
 def test_19_excel_ocr(self):
  from openpyxl import load_workbook
  j=self.execute('p_pdf_excel_ocr',[self.scan],lang='eng');values=[str(c.value) for sheet in load_workbook(j.output) for row in sheet for c in row];self.assertTrue(any('Apple' in value for value in values))
 def test_20_pdfa(self):
  j=self.execute('p_pdfa');report=json.loads(Path(j.output+'.report.json').read_text(encoding='utf-8'));self.assertEqual(report['part'],'2');self.assertEqual(report['conformance'].upper(),'B')
 def test_21_numbers(self):
  j=self.execute('p_numbers',pages='1,3',start=5);doc=fitz.open(j.output);self.assertIn('5 / 6',doc[0].get_text());self.assertNotIn('5 / 6',doc[1].get_text());doc.close()
 def test_22_watermark(self):
  j=self.execute('p_watermark',text='DRAFT');doc=fitz.open(j.output);self.assertIn('DRAFT',doc[0].get_text());doc.close()
 def test_23_crop(self):
  j=self.execute('p_crop',mode='margin',margin=10);doc=fitz.open(j.output);self.assertAlmostEqual(doc[0].rect.width,400-20*72/25.4,places=2);doc.close()
 def test_24_edit(self):
  j=self.execute('p_edit',actions=[dict(kind='text',page=1,rect=[.1,.6,.7,.1],normalized=True,text='Added text',size=12)]);doc=fitz.open(j.output);self.assertIn('Added text',doc[0].get_text());doc.close()
 def test_25_form_fill(self):
  j=self.execute('p_forms',mode='fill',actions=[dict(name='name',value='New name')]);doc=fitz.open(j.output);self.assertEqual(next(doc[0].widgets()).field_value,'New name');doc.close()
 def test_26_form_create(self):
  j=self.execute('p_forms',mode='create',actions=[dict(page=1,name='Accept',field_type='checkbox',rect=[.6,.6,.08,.06],normalized=True)]);doc=fitz.open(j.output);self.assertIn('Accept',[w.field_name for w in doc[0].widgets()]);doc.close()
 def test_27_sign_drawing(self):
  j=self.execute('p_sign',actions=[dict(kind='draw',page=1,rect=[.1,.6,.3,.1],normalized=True,points=[[.1,.6],[.2,.68],[.4,.62]],color='#000000')]);doc=fitz.open(j.output);self.assertTrue(doc[0].get_drawings());doc.close()
 def test_28_redact(self):
  j=self.execute('p_redact',text='SECRET',confirmed=True);doc=fitz.open(j.output);self.assertNotIn('SECRET',doc[0].get_text());self.assertNotIn('SECRET',str(doc.metadata));self.assertEqual(doc.embfile_count(),0);doc.close()
 def test_29_redact_confirmation(self):
  with self.assertRaises(ValueError):self.execute('p_redact',text='SECRET',confirmed=False)
 def test_30_compare(self):
  j=self.execute('p_compare',[self.source,self.source]);report=json.loads((Path(j.output)/'比較報告.json').read_text(encoding='utf-8'));self.assertFalse(any(r['changed'] for r in report))
 def test_31_markdown(self):
  j=self.execute('p_markdown');self.assertIn('Public information',(Path(j.output)/'文件.md').read_text(encoding='utf-8'))
 def test_32_content(self):
  j=self.execute('p_content');self.assertIn('CONFIDENTIAL',(Path(j.output)/'文字.txt').read_text(encoding='utf-8-sig'));self.assertTrue(list(Path(j.output).glob('附件*')))
 def test_33_flatten(self):
  j=self.execute('p_flatten');doc=fitz.open(j.output);self.assertFalse(list(doc[0].widgets() or []));doc.close()
 def test_34_copy(self):
  j=self.execute('p_copy');self.assertEqual(self.source.read_bytes(),Path(j.output).read_bytes())
 def test_35_summary_extract(self):
  j=self.execute('p_summary',mode='extract');self.assertTrue(list(Path(j.output).glob('*.pdf')))
 def test_36_preview(self):
  result=P.preview(str(self.source),2);self.assertEqual(result['count'],3);self.assertEqual(result['page'],2);self.assertTrue(result['src'].startswith('data:image/png'))
 def test_37_password_fail(self):
  encrypted=self.root/'encrypted.pdf';doc=fitz.open(self.source);doc.save(encrypted,encryption=fitz.PDF_ENCRYPT_AES_256,user_pw='abc',owner_pw='owner');doc.close()
  with self.assertRaises(ValueError):P.preview(str(encrypted),1,'wrong')
  self.assertEqual(P.preview(str(encrypted),1,'abc')['count'],3)
 def test_38_sanitized_secrets(self):
  self.assertNotIn('pfx_pw',W.clean_settings({'pfx_pw':'secret'}));self.assertEqual(E.public_options({'pfx_pw':'secret','actions':[1]}),{'pfx_pw':'[已隱藏]','actions':'[已隱藏]'})
 def test_39_paths(self):
  self.assertEqual(FA.completed_path('ZDGCFILE '+json.dumps(str(self.source)),str(self.root)),str(self.source));self.assertIsNone(FA.completed_path('ZDGCFILE NA',str(self.root)))
 def test_40_open_missing(self):
  with self.assertRaises(FileNotFoundError):FA.open_file(str(self.root/'missing.mp3'))
 def test_41_reveal(self):
  with patch('file_actions.subprocess.Popen') as process:FA.reveal(str(self.source));self.assertEqual(process.call_args[0][0][1],'/select,'+str(self.source))
 def test_42_open(self):
  with patch('file_actions.os.startfile') as launch:FA.open_file(json.dumps(str(self.source)));launch.assert_called_once_with(str(self.source))
 def test_43_certificate(self):
  from cryptography import x509
  from cryptography.x509.oid import NameOID
  from cryptography.hazmat.primitives import hashes,serialization
  from cryptography.hazmat.primitives.asymmetric import rsa
  from cryptography.hazmat.primitives.serialization import pkcs12
  import datetime
  key=rsa.generate_private_key(public_exponent=65537,key_size=2048);name=x509.Name([x509.NameAttribute(NameOID.COMMON_NAME,'Offline Test')]);now=datetime.datetime.now(datetime.timezone.utc)
  cert=x509.CertificateBuilder().subject_name(name).issuer_name(name).public_key(key.public_key()).serial_number(123).not_valid_before(now-datetime.timedelta(days=1)).not_valid_after(now+datetime.timedelta(days=1)).sign(key,hashes.SHA256())
  pfx=self.root/'test.pfx';pfx.write_bytes(pkcs12.serialize_key_and_certificates(b'test',key,cert,None,serialization.BestAvailableEncryption(b'test')))
  j=self.execute('p_certificate',pfx=str(pfx),pfx_pw='test')
  from pyhanko.pdf_utils.reader import PdfFileReader
  with open(j.output,'rb') as f:self.assertEqual(len(PdfFileReader(f).embedded_signatures),1)
 def test_44_cancel(self):
  j=E.Job([str(self.source)],'p_compress',{},tool='p_compress');j.cancelled=True
  with self.assertRaises(InterruptedError):P.run(self.eng,j)
 def test_45_offline_ai(self):
  j=self.execute('p_summary',mode='ai',language='English',length=2);self.assertTrue(list(Path(j.output).glob('*.md')))
 def test_46_translation(self):
  j=self.execute('p_translate',files=[self.source],language='English');self.assertTrue(list(Path(j.output).glob('*.pdf')))

if __name__=='__main__':unittest.main(verbosity=2)
