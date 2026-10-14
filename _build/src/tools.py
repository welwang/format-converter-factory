"""四維八德 tool box: video / audio / image / PDF / disc tools (no UI).

Every tool is described once in TOOLS (the UI builds its form from that) and implemented
by a function in RUN that receives (engine, job).  Jobs run inside Engine's thread pools,
so progress, cancel and the queue work exactly like normal conversions.
"""
import os, re, io, json, math, time, shutil, subprocess, tempfile, threading
import engine as E
import pdf_schema as PDF_SCHEMA

POS = [['br', '右下'], ['bl', '左下'], ['tr', '右上'], ['tl', '左上'], ['c', '正中央'], ['b', '下方置中'], ['t', '上方置中']]
RATIO = [['9:16', '9:16 直式短影音'], ['1:1', '1:1 方形'], ['4:5', '4:5 IG 直式'], ['16:9', '16:9 橫式'], ['4:3', '4:3'], ['21:9', '21:9 電影寬螢幕']]
LANGS = [['auto', '自動偵測'], ['zh', '中文'], ['en', '英文'], ['ja', '日文'], ['ko', '韓文'], ['yue', '粵語'], ['th', '泰文'], ['vi', '越南文'], ['id', '印尼文'], ['fr', '法文'], ['de', '德文'], ['es', '西班牙文']]
AUDIO_OUT = [['same', '跟原檔相同'], ['mp3', 'MP3'], ['m4a', 'M4A'], ['flac', 'FLAC'], ['wav', 'WAV'], ['opus', 'Opus']]
IMG_OUT = [['same', '跟原檔相同'], ['jpg', 'JPG'], ['png', 'PNG'], ['webp', 'WebP']]

def F(k, label, type='number', default=None, **kw):
    d = dict(k=k, label=label, type=type, default=default); d.update(kw); return d

# ---------------------------------------------------------------- registry
TOOLS = [
    # ---------------- video
    dict(id='v_social', group='video', icon='phone', name='社群一鍵規格', desc='直式短影音、方形、壓到 10／20 MB，一鍵搞定', accept=['video'], fields=[
        F('mode', '用途', 'select', 'vertical', options=[['vertical', '直式短影音 1080×1920（模糊背景填滿）'], ['square', '方形 1080×1080（模糊背景填滿）'],
                                                      ['mb10', '壓到 10 MB 以內（Discord 等）'], ['mb20', '壓到 20 MB 以內（Email 附件）'], ['hd720', '720p 好分享（畫質與大小平衡）']])]),
    dict(id='v_fill', group='video', icon='aspect', name='改比例（模糊背景）', desc='不裁掉任何畫面，空白處用模糊背景補滿', accept=['video'], fields=[
        F('ratio', '目標比例', 'select', '9:16', options=RATIO), F('height', '輸出高度', 'select', '1920', options=[['1920', '1920'], ['1080', '1080'], ['1280', '1280'], ['720', '720']]),
        F('blur', '背景模糊度', 'number', 20, min=2, max=60)]),
    dict(id='v_crop', group='video', icon='crop', name='裁切畫面', desc='依比例置中裁切，或在畫面上框選範圍', accept=['video'], fields=[
        F('mode', '方式', 'select', '9:16', options=RATIO + [['custom', '自己框選']]), F('rect', '框選範圍', 'rect', '', show={'mode': 'custom'})]),
    dict(id='v_rotate', group='video', icon='rotate', name='旋轉／翻轉', desc='轉正手機直拍影片、左右鏡像', accept=['video'], fields=[
        F('op', '動作', 'select', 'cw', options=[['cw', '順時針 90°'], ['ccw', '逆時針 90°'], ['180', '旋轉 180°'], ['hflip', '左右翻轉'], ['vflip', '上下翻轉']]),
        F('fast', '只改方向標記（瞬間完成、不重新編碼；少數舊軟體看不到效果）', 'check', False, show={'op': ['cw', 'ccw', '180']})]),
    dict(id='v_speed', group='video', icon='speed', name='變速', desc='快轉、慢動作，聲音音調不變', accept=['video'], fields=[
        F('speed', '速度（倍）', 'number', 2, min=0.1, max=16, step=0.05), F('keep_audio', '保留聲音', 'check', True)]),
    dict(id='v_wm_text', group='video', icon='textmark', name='文字浮水印', desc='加上署名、網址或版權文字', accept=['video'], fields=[
        F('text', '文字', 'text', '© 我的頻道'), F('pos', '位置', 'select', 'br', options=POS), F('size', '字級（畫面高度 %）', 'number', 5, min=1, max=30),
        F('color', '顏色', 'color', '#ffffff'), F('opacity', '不透明度 %', 'number', 80, min=5, max=100), F('shadow', '加陰影（任何背景都看得清楚）', 'check', True)]),
    dict(id='v_wm_img', group='video', icon='logo', name='圖片浮水印／Logo', desc='把 Logo 疊到影片角落', accept=['video'], fields=[
        F('logo', 'Logo 圖片', 'file', '', filter='圖片 (*.png;*.jpg;*.webp)'), F('pos', '位置', 'select', 'br', options=POS),
        F('size', '寬度（畫面寬度 %）', 'number', 15, min=2, max=80), F('opacity', '不透明度 %', 'number', 90, min=5, max=100)]),
    dict(id='v_delogo', group='video', icon='eraser', name='去除浮水印', desc='框選台標或浮水印，自動用周圍畫面補掉', accept=['video'], fields=[
        F('rect', '浮水印位置', 'rect', '')]),
    dict(id='v_subs', group='video', icon='captions', name='加字幕', desc='把 SRT／ASS 字幕燒進畫面，或內嵌成可開關的字幕軌', accept=['video'], fields=[
        F('sub', '字幕檔', 'file', '', filter='字幕 (*.srt;*.ass;*.ssa;*.vtt)'), F('mode', '方式', 'select', 'burn', options=[['burn', '燒進畫面（任何播放器都看得到）'], ['soft', '內嵌字幕軌（可開關、不重新編碼）']]),
        F('size', '字級', 'number', 22, min=8, max=80, show={'mode': 'burn'}), F('outline', '黑色外框', 'check', True, show={'mode': 'burn'})]),
    dict(id='v_autosub', group='video', icon='sparkles', name='AI 自動字幕', desc='離線語音辨識產生字幕，可轉繁體、可直接燒進影片', accept=['video', 'audio'], fields=[
        F('lang', '語言', 'select', 'zh', options=LANGS), F('model', '辨識模型', 'select', 'base', options=[['base', '離線快速（模型已內建）']]),
        F('trad', '轉成繁體中文', 'check', True), F('out', '輸出', 'select', 'srt', options=[['srt', '只要 SRT 字幕檔'], ['burn', '字幕檔＋燒進影片'], ['soft', '字幕檔＋內嵌字幕軌']])]),
    dict(id='v_audio', group='video', icon='music', name='配樂／換音軌', desc='加背景音樂、與原聲混合，或整個換掉', accept=['video'], fields=[
        F('music', '音樂檔', 'file', '', filter='音訊 (*.mp3;*.m4a;*.wav;*.flac;*.ogg;*.aac)'), F('mode', '方式', 'select', 'mix', options=[['mix', '與原聲混合'], ['replace', '取代原聲']]),
        F('music_vol', '音樂音量 %', 'number', 35, min=0, max=200), F('orig_vol', '原聲音量 %', 'number', 100, min=0, max=200, show={'mode': 'mix'}),
        F('loop', '音樂太短時循環播放', 'check', True), F('fade', '結尾淡出（秒）', 'number', 2, min=0, max=10)]),
    dict(id='v_mute', group='video', icon='mute', name='去除聲音', desc='只留畫面，瞬間完成', accept=['video'], fields=[]),
    dict(id='v_split', group='video', icon='scissors', name='分割影片', desc='每隔幾秒切一段，或平均切成幾段', accept=['video', 'audio'], fields=[
        F('mode', '方式', 'select', 'every', options=[['every', '每隔固定秒數'], ['parts', '平均切成幾段']]), F('sec', '秒數', 'number', 60, min=1, show={'mode': 'every'}),
        F('parts', '段數', 'number', 3, min=2, max=100, show={'mode': 'parts'}), F('exact', '精準切割（重新編碼，較慢；不勾則在關鍵影格切、瞬間完成）', 'check', False)]),
    dict(id='v_stab', group='video', icon='stabilize', name='防手震', desc='分析晃動後穩定畫面（邊緣會略為放大）', accept=['video'], fields=[
        F('strength', '強度', 'select', '15', options=[['8', '輕微'], ['15', '一般'], ['30', '強力']])]),
    dict(id='v_enhance', group='video', icon='sliders', name='畫面調整', desc='亮度、對比、飽和度、降噪、銳利化', accept=['video'], fields=[
        F('bright', '亮度', 'number', 0, min=-100, max=100), F('contrast', '對比', 'number', 0, min=-100, max=100), F('sat', '飽和度', 'number', 0, min=-100, max=100),
        F('denoise', '降噪', 'select', '0', options=[['0', '不降噪'], ['1', '輕微'], ['2', '強力']]), F('sharpen', '銳利化', 'check', False)]),
    dict(id='v_fade', group='video', icon='fade', name='淡入淡出', desc='片頭從黑畫面淡入、片尾淡出', accept=['video'], fields=[
        F('fin', '淡入（秒）', 'number', 1, min=0, max=10, step=0.5), F('fout', '淡出（秒）', 'number', 1, min=0, max=10, step=0.5)]),
    dict(id='v_reverse', group='video', icon='rewind', name='倒轉播放', desc='影片倒著播（適合 3 分鐘內的短片）', accept=['video'], fields=[
        F('audio', '聲音也倒轉', 'check', True)]),
    dict(id='v_frames', group='video', icon='frames', name='擷取畫面／縮圖總覽', desc='每隔幾秒存一張圖，或做成一張縮圖總覽', accept=['video'], fields=[
        F('mode', '方式', 'select', 'sheet', options=[['sheet', '一張縮圖總覽'], ['every', '每隔幾秒存一張']]), F('sec', '間隔（秒）', 'number', 5, min=0.2, show={'mode': 'every'}),
        F('cols', '欄數', 'number', 4, min=1, max=10, show={'mode': 'sheet'}), F('rows', '列數', 'number', 5, min=1, max=20, show={'mode': 'sheet'})]),
    dict(id='v_grid', group='video', icon='grid', name='多支影片並排', desc='2～4 支影片左右、上下或田字排列成一支', accept=['video'], together=True, fields=[
        F('layout', '排列', 'select', 'h', options=[['h', '左右並排'], ['v', '上下並排'], ['grid', '田字（4 支）']]), F('size', '每格高度', 'select', '720', options=[['1080', '1080'], ['720', '720'], ['480', '480']])]),
    # ---------------- audio
    dict(id='a_volume', group='audio', icon='volume', name='音量／淡入淡出', desc='調大調小、標準化響度、頭尾淡入淡出（影片也可用）', accept=['audio', 'video'], fields=[
        F('gain', '音量增減（dB）', 'number', 0, min=-30, max=30), F('norm', '響度標準化（-16 LUFS）', 'check', False),
        F('fin', '淡入（秒）', 'number', 0, min=0, max=30), F('fout', '淡出（秒）', 'number', 0, min=0, max=30), F('fmt', '輸出格式', 'select', 'same', options=AUDIO_OUT)]),
    dict(id='a_pitch', group='audio', icon='pitch', name='變速／變調', desc='升降 Key、加快放慢，互不影響', accept=['audio', 'video'], fields=[
        F('tempo', '速度 %', 'number', 100, min=25, max=400), F('semi', '升降半音（Key）', 'number', 0, min=-12, max=12), F('fmt', '輸出格式', 'select', 'same', options=AUDIO_OUT)]),
    dict(id='a_vocal', group='audio', icon='mic', name='去人聲（伴奏）', desc='消去置中的人聲，做卡拉 OK 伴奏（需立體聲）', accept=['audio', 'video'], fields=[
        F('bass', '保留低音', 'check', True), F('fmt', '輸出格式', 'select', 'same', options=AUDIO_OUT)]),
    dict(id='a_denoise', group='audio', icon='wind', name='降噪', desc='去除底噪、風切、冷氣聲（影片也可用）', accept=['audio', 'video'], fields=[
        F('level', '強度', 'select', '12', options=[['6', '輕微'], ['12', '一般'], ['24', '強力']]), F('fmt', '輸出格式', 'select', 'same', options=AUDIO_OUT)]),
    dict(id='a_silence', group='audio', icon='silence', name='去除靜音', desc='自動剪掉錄音中的空白停頓', accept=['audio'], fields=[
        F('db', '靜音門檻（dB）', 'number', -45, min=-80, max=-10), F('min', '超過幾秒才剪', 'number', 0.8, min=0.1, max=10, step=0.1), F('fmt', '輸出格式', 'select', 'same', options=AUDIO_OUT)]),
    dict(id='a_mix', group='audio', icon='mix', name='混音', desc='多個音軌疊在一起同時播放', accept=['audio', 'video'], together=True, fields=[
        F('fmt', '輸出格式', 'select', 'mp3', options=AUDIO_OUT[1:]), F('dur', '長度', 'select', 'longest', options=[['longest', '以最長的為準'], ['first', '以第一個為準'], ['shortest', '以最短的為準']])]),
    dict(id='a_ringtone', group='audio', icon='bell', name='鈴聲製作', desc='擷取一段做成 iPhone（M4R）或 Android 鈴聲', accept=['audio', 'video'], fields=[
        F('start', '從第幾秒開始', 'number', 0, min=0), F('len', '長度（秒）', 'number', 30, min=3, max=40), F('fmt', '格式', 'select', 'm4r', options=[['m4r', 'iPhone（M4R）'], ['mp3', 'Android（MP3）']])]),
    dict(id='a_tags', group='audio', icon='tag', name='歌曲資訊／封面', desc='修改歌名、歌手、專輯，加上封面圖（不重新編碼）', accept=['audio'], fields=[
        F('title', '歌名', 'text', ''), F('artist', '歌手', 'text', ''), F('album', '專輯', 'text', ''), F('year', '年份', 'text', ''), F('cover', '封面圖片', 'file', '', filter='圖片 (*.jpg;*.png)')]),
    dict(id='a_transcribe', group='audio', icon='captions', name='音檔轉文字', desc='錄音、影片離線轉成逐字稿（TXT／Word／SRT），可轉繁體', accept=['audio', 'video'], fields=[
        F('lang', '語言', 'select', 'zh', options=LANGS), F('model', '辨識模型', 'select', 'base', options=[['base', '離線快速（已內建）']]),
        F('trad', '轉成繁體中文', 'check', True), F('fmt', '輸出', 'select', 'docx', options=[['docx', 'Word 逐字稿（含時間）'], ['txt', '純文字'], ['txt_ts', '純文字（每段附時間）'], ['srt', 'SRT 字幕']])]),
    dict(id='a_meeting', group='audio', icon='list', name='會議總結', desc='會議錄音離線轉逐字稿，整理出摘要、重點、決議與待辦事項', accept=['audio', 'video'], fields=[
        F('title', '會議名稱（選填）', 'text', ''), F('lang', '語言', 'select', 'zh', options=LANGS), F('model', '辨識模型', 'select', 'base', options=[['base', '離線快速（已內建）']]),
        F('mode', '整理方式', 'select', 'ai', options=[['ai', '離線 AI 整理（較慢、較完整）'], ['extract', '快速摘錄（不使用 AI 模型）']]),
        F('points', '重點數量', 'number', 8, min=3, max=20)]),
    # ---------------- image
    dict(id='i_resize', group='image', icon='resize', name='批次縮放', desc='依最長邊、寬、高或百分比縮放，不會放大', accept=['image'], fields=[
        F('mode', '依據', 'select', 'long', options=[['long', '最長邊（像素）'], ['w', '寬度（像素）'], ['h', '高度（像素）'], ['pct', '百分比']]), F('val', '數值', 'number', 1920, min=1),
        F('fmt', '輸出格式', 'select', 'same', options=IMG_OUT)]),
    dict(id='i_rotate', group='image', icon='rotate', name='旋轉／翻轉／裁切', desc='批次轉正、鏡像、依比例裁切', accept=['image'], fields=[
        F('op', '旋轉', 'select', 'none', options=[['none', '不旋轉'], ['cw', '順時針 90°'], ['ccw', '逆時針 90°'], ['180', '180°'], ['hflip', '左右翻轉'], ['vflip', '上下翻轉']]),
        F('crop', '裁切比例', 'select', 'none', options=[['none', '不裁切']] + RATIO + [['custom', '自己框選']]), F('rect', '框選範圍', 'rect', '', show={'crop': 'custom'}),
        F('fmt', '輸出格式', 'select', 'same', options=IMG_OUT)]),
    dict(id='i_watermark', group='image', icon='copyright', name='圖片浮水印', desc='文字或 Logo，可斜向平鋪防盜用', accept=['image'], fields=[
        F('kind', '類型', 'select', 'text', options=[['text', '文字'], ['logo', '圖片 Logo']]), F('text', '文字', 'text', '© 版權所有', show={'kind': 'text'}),
        F('logo', 'Logo 圖片', 'file', '', filter='圖片 (*.png;*.jpg;*.webp)', show={'kind': 'logo'}), F('pos', '位置', 'select', 'br', options=POS + [['tile', '斜向平鋪（防盜）']]),
        F('size', '大小（圖片寬度 %）', 'number', 20, min=2, max=90), F('color', '文字顏色', 'color', '#ffffff', show={'kind': 'text'}), F('opacity', '不透明度 %', 'number', 70, min=5, max=100),
        F('fmt', '輸出格式', 'select', 'same', options=IMG_OUT)]),
    dict(id='i_collage', group='image', icon='collage', name='拼長圖／拼圖', desc='多張圖片接成長圖、橫圖或九宮格', accept=['image'], together=True, fields=[
        F('layout', '排列', 'select', 'v', options=[['v', '直向長圖'], ['h', '橫向拼接'], ['grid', '格狀（九宮格）']]), F('cols', '每列張數', 'number', 3, min=2, max=10, show={'layout': 'grid'}),
        F('size', '統一寬度／高度（像素）', 'number', 1080, min=100, max=8000), F('gap', '間距（像素）', 'number', 0, min=0, max=200), F('bg', '背景色', 'color', '#ffffff'),
        F('fmt', '輸出格式', 'select', 'jpg', options=IMG_OUT[1:])]),
    dict(id='i_slideshow', group='image', icon='slideshow', name='圖片轉影片', desc='照片做成有轉場的 MP4，可加背景音樂', accept=['image'], together=True, fields=[
        F('sec', '每張秒數', 'number', 3, min=0.5, max=60, step=0.5), F('trans', '轉場', 'select', 'fade', options=[['fade', '淡化'], ['slideleft', '左滑'], ['circleopen', '圓形展開'], ['dissolve', '溶解'], ['none', '無']]),
        F('tdur', '轉場秒數', 'number', 0.6, min=0.1, max=3, step=0.1), F('size', '解析度', 'select', '1920x1080', options=[['1920x1080', '1920×1080 橫式'], ['1080x1920', '1080×1920 直式'], ['1080x1080', '1080×1080 方形'], ['1280x720', '1280×720']]),
        F('music', '背景音樂（選填）', 'file', '', filter='音訊 (*.mp3;*.m4a;*.wav;*.flac)')]),
    dict(id='i_gif', group='image', icon='gif', name='圖片轉 GIF 動畫', desc='多張圖片做成會動的 GIF', accept=['image'], together=True, fields=[
        F('ms', '每張毫秒', 'number', 500, min=20, max=10000), F('w', '寬度（像素）', 'number', 600, min=50, max=2000), F('loop', '無限循環', 'check', True)]),
    # ---------------- pdf
    dict(id='p_split', group='pdf', icon='pages', name='PDF 分割', desc='每頁一個檔，或依頁碼範圍拆開', accept=['pdf'], fields=[
        F('mode', '方式', 'select', 'each', options=[['each', '每頁一個檔案'], ['ranges', '依範圍（例如 1-3, 4-10）'], ['every', '每幾頁一個檔案']]),
        F('ranges', '範圍', 'text', '1-3, 4-', show={'mode': 'ranges'}), F('n', '頁數', 'number', 5, min=1, show={'mode': 'every'})]),
    dict(id='p_pages', group='pdf', icon='list', name='頁面整理', desc='挑出、刪除或重新排列頁面', accept=['pdf'], fields=[
        F('mode', '動作', 'select', 'keep', options=[['keep', '只保留（依輸入順序，可重排）'], ['delete', '刪除這些頁']]), F('pages', '頁碼', 'text', '1, 3-5', placeholder='例如 1, 3-5, 8-')]),
    dict(id='p_rotate', group='pdf', icon='rotate', name='PDF 旋轉', desc='轉正掃描歪掉的頁面', accept=['pdf'], fields=[
        F('angle', '角度', 'select', '90', options=[['90', '順時針 90°'], ['270', '逆時針 90°'], ['180', '180°']]), F('pages', '頁碼（空白＝全部）', 'text', '')]),
    dict(id='p_encrypt', group='pdf', icon='lock', name='PDF 加密', desc='設定開啟密碼', accept=['pdf'], fields=[
        F('pw', '密碼', 'text', ''), F('print', '允許列印', 'check', True), F('copy', '允許複製文字', 'check', False)]),
    dict(id='p_decrypt', group='pdf', icon='unlock', name='PDF 解除密碼', desc='輸入密碼後存成沒有密碼的版本', accept=['pdf'], fields=[F('pw', '密碼', 'text', '')]),
    # ---------------- any
    dict(id='x_privacy', group='other', icon='shield', name='隱私清除', desc='移除照片 GPS／拍攝資訊、影片位置、PDF 作者（JPG／PNG 無損）', accept=['image', 'video', 'audio', 'pdf'], fields=[]),
    # ---------------- disc
    dict(id='d_dvd', group='disc', icon='disc', name='DVD 轉影片', desc='光碟或 VIDEO_TS 資料夾轉成 MP4（無防拷光碟）', accept=['disc'], fields=[
        F('deint', '去交錯（DVD 多為交錯掃描，建議開）', 'check', True), F('fmt', '輸出', 'select', 'mp4', options=[['mp4', 'MP4'], ['mkv', 'MKV（保留所有音軌與字幕）']])]),
    dict(id='d_bluray', group='disc', icon='disc', name='藍光轉影片', desc='無防拷藍光光碟或資料夾轉 MP4／MKV', accept=['disc'], fields=[
        F('fmt', '輸出', 'select', 'mp4', options=[['mp4', 'MP4'], ['mkv', 'MKV']])]),
    dict(id='d_cd', group='disc', icon='cd', name='音樂 CD 轉檔', desc='CD 音軌轉 MP3／FLAC', accept=['disc'], fields=[
        F('fmt', '格式', 'select', 'flac', options=[['flac', 'FLAC（無損）'], ['mp3', 'MP3'], ['m4a', 'M4A'], ['wav', 'WAV']])]),
    dict(id='d_iso', group='disc', icon='archive', name='光碟轉 ISO', desc='資料光碟、DVD 備份成 ISO 映像檔', accept=['disc'], fields=[]),
]
TOOLS += PDF_SCHEMA.TOOLS
TOOL = {t['id']: t for t in TOOLS}
GROUPS = [['video', '影片工具'], ['audio', '音訊工具'], ['image', '圖片工具'], ['pdf', 'PDF 工具'], ['other', '隱私工具'], ['disc', '光碟']]

MODEL_LABEL = {'base': '離線快速（base）', 'small': '離線準確（small）', 'medium': '離線高準確（medium，需 8 GB 以上）', 'turbo': '離線最準確（large-v3-turbo）'}

def installed_models():
    return [k for k in MODELS if find_model(MODELS[k][0], MODELS[k][1] * 0.98)]

def schema():
    import pdf_tools
    models = installed_models()
    opts = [[k, MODEL_LABEL[k]] for k in models] or [['base', '缺少離線辨識模型']]
    tools = []
    for t in TOOLS:
        if t['id'] in ('v_autosub', 'a_transcribe', 'a_meeting'):
            t = dict(t, fields=[dict(f, options=opts, default=opts[0][0]) if f['k'] == 'model' else f for f in t['fields']])
        tools.append(t)
    return dict(groups=GROUPS, tools=tools, font=bool(find_font()), whisper=whisper_ok() and bool(models), whisper_ffmpeg=whisper_ok(),
                speech_models=models, ai=ai_ok(), pdf=pdf_tools.capabilities())

def ai_ok():
    return os.path.isfile(E.bundled('pdf-ai', 'llama-completion.exe')) and os.path.isfile(E.bundled('pdf-ai', 'model.gguf'))

def run(eng, j):
    fn = RUN.get(j.tool)
    if not fn: raise RuntimeError('未知的工具')
    fn(eng, j, j.opt)

# ---------------------------------------------------------------- helpers
FF = lambda: E.FFMPEG

def need_ff():
    if not E.FFMPEG: raise RuntimeError('找不到 ffmpeg（tools 資料夾）')

def short_path(p):
    """8.3 short path on Windows (ASCII, no quotes/brackets) - the safest way to hand odd paths to an ffmpeg filter."""
    if os.name != 'nt' or not os.path.exists(p): return p
    import ctypes
    buf = ctypes.create_unicode_buffer(1024)
    n = ctypes.windll.kernel32.GetShortPathNameW(str(p), buf, 1024)
    return buf.value if 0 < n < 1024 else p

def ffesc(p):
    """Escape a path for use inside a single-quoted ffmpeg filter option value."""
    p = str(p)
    if os.name == 'nt': p = short_path(p)      # ASCII-only when 8.3 names exist; harmless otherwise
    if "'" in p: p = p.replace("'", '')          # cannot appear inside ffmpeg single quotes; short path normally removes it
    return p.replace('\\', '/').replace(':', r'\:')

def num(o, k, d=0.0):
    try: return float(o.get(k) if o.get(k) not in (None, '') else d)
    except (TypeError, ValueError): return float(d)

def out_for(j, src, ext, suffix, folder=None):
    folder = folder or j.opt.get('out_dir') or (os.path.dirname(src) if os.path.isfile(src) else default_dir(j))
    os.makedirs(folder, exist_ok=True)
    if re.fullmatch(r'[A-Za-z]:[\\/]?', src): stem = '光碟'
    else: stem = os.path.splitext(os.path.basename(src.rstrip('\\/')))[0] or 'output'
    return E.unique_path(os.path.join(folder, f'{stem}{suffix}.{ext}'), avoid=j.files)

def out_dir_for(j, src, suffix):
    p = out_for(j, src, 'x', suffix)[:-2]
    p = E.unique_path(p); os.makedirs(p); return p

def default_dir(j):
    sub = 'Music' if j.tool in ('d_cd',) else 'Videos'
    d = os.path.join(os.path.expanduser('~'), sub)
    return d if os.path.isdir(d) else os.path.expanduser('~')

def vext(src):
    e = os.path.splitext(src)[1].lower().lstrip('.')
    return e if e in ('mp4', 'mkv', 'mov', 'webm') else 'mp4'

def aext(src, o):
    f = o.get('fmt') or 'same'
    if f != 'same': return f
    e = os.path.splitext(src)[1].lower().lstrip('.')
    return e if e in ('mp3', 'm4a', 'flac', 'wav', 'opus', 'ogg', 'aac') else 'mp3'

def info(src): return E.probe_info(src)
def has_audio(src): return any(t == 'audio' for t, _ in info(src)['streams'])
def has_video(src): return any(t == 'video' and c not in ('mjpeg', 'png') for t, c in info(src)['streams'])

def base_cmd(extra_in=()):
    return [FF(), '-hide_banner', '-nostdin', '-y', '-progress', 'pipe:1', '-nostats', '-v', 'error'] + list(extra_in)

_NOISE = ('Nothing was written', 'Terminating thread', 'Task finished with error', 'Could not open encoder before EOF', 'Conversion failed',
          'Error while filtering', 'Error sending frames', 'Error marking filters', 'Error closing', 'Error while processing the decoded data')
def ff_error(err, default='ffmpeg 失敗'):
    """Pick the line that actually explains an ffmpeg failure (the last line is usually a generic follow-up)."""
    lines = [re.sub(r'^(\[[^\]]*\]\s*)+', '', l).strip() for l in (err or '').splitlines() if l.strip()]
    good = [l for l in lines if l and not any(n in l for n in _NOISE)]
    msg = ((good[0] if good else (lines[-1] if lines else default)) or default)
    for k, zh in (('No space left', '磁碟空間不足'), ('Permission denied', '沒有寫入權限（檔案可能正被其他程式開啟）'),
                  ('Invalid data found', '檔案損毀或格式無法辨識'), ('No such file', '找不到檔案'), ('does not contain any stream', '檔案裡沒有可用的影音內容'),
                  ('Stream map', '找不到需要的影像或聲音軌'), ('out of memory', '記憶體不足'), ('Cannot allocate memory', '記憶體不足')):
        if k.lower() in msg.lower(): return f'{zh}（{msg[:160]}）'
    return msg[:250]

def ff_run(eng, j, build, dur, cwd=None, base=0.0, span=1.0, gpu=True):
    """build(hw_list) -> cmd.  Runs ffmpeg; if a GPU encoder fails, rebuilds the command for the CPU and retries."""
    hw = E.detect_hw() if (gpu and j.opt.get('hwaccel', True)) else []
    cmd = build(hw)
    eng._exec(j, cmd, dur, cwd=cwd, base=base, span=span)
    if j._rc != 0 and hw and not j.cancelled and any(h in cmd for h in hw):
        cmd = build([]); eng._exec(j, cmd, dur, cwd=cwd, base=base, span=span)
    if j._rc != 0 and not j.cancelled:
        raise RuntimeError(ff_error(j._err))

def venc_args(j, out, dur, hw, audio=True):
    """Encoder arguments for a video tool output (uses the global quality preset / GPU)."""
    t = os.path.splitext(out)[1].lower().lstrip('.')
    o = dict(j.opt); o.pop('target_mb', None)
    va, enc = E.video_args(t if t in ('mp4', 'mkv', 'mov', 'webm') else 'mp4', o, dur, hw)
    j.encoder = enc or ''
    if not audio:
        k = va.index('-c:a'); va = va[:k] + ['-an'] + va[k + 4:]
    return va

def video_tool(eng, j, src, suffix, vf=None, af=None, fc=None, maps=None, extra_in=(), pre_in=(), dur=None, audio=True, cwd=None, ext=None, base=0.0, span=1.0):
    need_ff()
    out = out_for(j, src, ext or vext(src), suffix); j.output = out
    d = dur if dur is not None else info(src)['dur']
    audio = audio and has_audio(src)
    tm = E.hdr_filters(src)
    if tm:                                   # HDR source: tone-map to SDR before any other filter
        if fc: fc = fc.replace('[0:v]', '[0:v]' + ','.join(tm) + ',', 1)
        else: vf = tm + list(vf or [])
        j.hdr = True
    def build(hw):
        c = base_cmd(pre_in) + ['-i', src]
        for x in extra_in: c += list(x)
        if fc:
            c += ['-filter_complex', fc]
            for m in maps: c += ['-map', m]
        else:
            if vf: c += ['-vf', ','.join(vf)]
            if af and audio: c += ['-af', ','.join(af)]
            c += ['-map', '0:v:0'] + (['-map', '0:a:0'] if audio else [])
        c += venc_args(j, out, d, hw, audio=audio or bool(fc and any('a' in m for m in maps)))
        return c + ['-map_metadata', '0', out]
    ff_run(eng, j, build, d, cwd=cwd, base=base, span=span)
    return out

def audio_tool(eng, j, src, suffix, af, ext=None, extra_in=(), fc=None, maps=None, dur=None, pre_in=()):
    """Audio filter on an audio file, or on the audio of a video (picture copied untouched)."""
    need_ff()
    isvid = has_video(src) and (ext or 'same') == 'same' and j.opt.get('fmt', 'same') == 'same'
    if isvid:
        out = out_for(j, src, vext(src), suffix)
    else:
        out = out_for(j, src, ext or aext(src, j.opt), suffix)
    j.output = out
    d = dur if dur is not None else info(src)['dur']
    t = os.path.splitext(out)[1].lstrip('.').lower()
    c = base_cmd(pre_in) + ['-i', src]
    for x in extra_in: c += list(x)
    if fc:
        c += ['-filter_complex', fc] + sum((['-map', m] for m in maps), [])
    else:
        c += ['-map', '0:a:0'] + (['-af', ','.join(af)] if af else [])
    if isvid:
        c += ['-map', '0:v:0', '-c:v', 'copy', '-c:a', 'aac', '-b:a', f"{E.PRESET[j.opt.get('preset', 'balanced')]['aac']}k"]
        if t in ('mp4', 'mov'): c += ['-movflags', '+faststart']
        j.encoder = 'FFmpeg（畫面不重新編碼）'
    else:
        c += ['-vn'] + (E.audio_args(t, j.opt) if t in E.AUDIO_COPY or t == 'wav' else ['-c:a', 'aac', '-b:a', '192k'])
        if t == 'm4r': c += ['-f', 'ipod']
        j.encoder = 'FFmpeg'
    c += ['-map_metadata', '0', out]
    eng._exec(j, c, d)
    if j._rc != 0 and not j.cancelled:
        raise RuntimeError(ff_error(j._err))
    return out

def find_font(bold=False):
    """Windows' Microsoft JhengHei when present, otherwise the Noto Sans TC shipped in tools\\fonts
    (so watermarks / subtitles work the same on any PC, including non-Chinese Windows)."""
    win = os.environ.get('WINDIR', r'C:\Windows')
    cands = [os.path.join(win, 'Fonts', n) for n in (('msjhbd.ttc', 'msjh.ttc') if bold else ('msjh.ttc', 'msjhbd.ttc'))]
    cands += [E.bundled('fonts', 'NotoSansTC-Bold.otf' if bold else 'NotoSansTC-Regular.otf'), E.bundled('fonts', 'NotoSansTC-Regular.otf')]
    cands += ['/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
    for p in cands:
        if os.path.isfile(p): return p
    return None

def font_family():
    f = os.path.basename(find_font() or '').lower()
    return 'Microsoft JhengHei' if f.startswith('msjh') else 'Noto Sans TC' if f.startswith('notosanstc') else 'Noto Sans CJK TC'

def hexcolor(c, default='#ffffff'):
    c = (c or default).strip()
    return c if re.fullmatch(r'#[0-9a-fA-F]{6}', c) else default

def pos_xy(pos, W='W', H='H', w='w', h='h', m='M'):
    """Overlay position expressions (main W/H, overlay w/h, margin m)."""
    x = {'l': m, 'c': f'({W}-{w})/2', 'r': f'{W}-{w}-{m}'}
    y = {'t': m, 'c': f'({H}-{h})/2', 'b': f'{H}-{h}-{m}'}
    return {'br': (x['r'], y['b']), 'bl': (x['l'], y['b']), 'tr': (x['r'], y['t']), 'tl': (x['l'], y['t']),
            'c': (x['c'], y['c']), 'b': (x['c'], y['b']), 't': (x['c'], y['t'])}.get(pos, (x['r'], y['b']))

def parse_rect(s):
    try:
        x, y, w, h = [int(round(float(v))) for v in str(s).split(',')]
        if w > 1 and h > 1: return x, y, w, h
    except Exception: pass
    raise RuntimeError('請先在畫面上框選範圍')

def ratio_of(r):
    a, b = r.split(':'); return float(a) / float(b)

def atempo_chain(f):
    out = []
    while f > 100: out.append('atempo=100'); f /= 100
    while f < 0.5: out.append('atempo=0.5'); f /= 0.5
    out.append(f'atempo={f:.4f}')
    return out

def tmpdir():
    return tempfile.mkdtemp(prefix='zdgc_t_')

# ---------------------------------------------------------------- video tools
def fill_graph(W, H, blur, audio):
    fc = (f'[0:v]split=2[a][b];[a]scale={W}:{H}:force_original_aspect_ratio=increase,crop={W}:{H},boxblur={blur}:2,eq=brightness=-0.06[bg];'
          f'[b]scale={W}:{H}:force_original_aspect_ratio=decrease[fg];[bg][fg]overlay=(W-w)/2:(H-h)/2,setsar=1,format=yuv420p[v]')
    return fc, ['[v]'] + (['0:a:0'] if audio else [])

def even(x): return int(x) // 2 * 2

def t_fill(eng, j, o, ratio=None, H=None, suffix=None):
    src = j.files[0]; r = ratio_of(ratio or o.get('ratio', '9:16')); H = int(H or o.get('height') or 1920)
    W = even(H * r); fc, maps = fill_graph(W, H, int(num(o, 'blur', 20)), has_audio(src))
    video_tool(eng, j, src, suffix or f'_{(ratio or o.get("ratio", "9:16")).replace(":", "x")}', fc=fc, maps=maps)

def t_social(eng, j, o):
    src = j.files[0]; m = o.get('mode', 'vertical')
    if m == 'vertical': return t_fill(eng, j, dict(o, blur=20), '9:16', 1920, '_直式')
    if m == 'square': return t_fill(eng, j, dict(o, blur=20), '1:1', 1080, '_方形')
    # size-limited / 720p: reuse the normal converter (target size, GPU, never-bigger)
    o2 = dict(o, fastcopy=False)
    if m == 'mb10': o2.update(target_mb=9.5, height=min(720, info(src)['height'] or 720))
    elif m == 'mb20': o2.update(target_mb=19, height=min(1080, info(src)['height'] or 1080))
    else: o2.update(height=min(720, info(src)['height'] or 720))
    if m in ('mb10', 'mb20') and info(src)['dur'] > 0:
        kbps = float(o2['target_mb']) * 8192 / info(src)['dur']
        if kbps < 700: o2['height'] = min(o2['height'], 480)
        if kbps < 350: o2['height'] = min(o2['height'], 360)
    j.opt = o2; j.target = 'mp4'; j.cat = 'video'
    suffix = {'mb10': '_10MB', 'mb20': '_20MB', 'hd720': '_720p'}[m]
    j.opt['_suffix'] = suffix
    try: eng._ffmpeg(j)
    finally: j.cat = 'tool'

def t_crop(eng, j, o):
    src = j.files[0]; m = o.get('mode', '9:16')
    if m == 'custom':
        x, y, w, h = parse_rect(o.get('rect'))
        vf = [f'crop={even(w)}:{even(h)}:{x}:{y}']
    else:
        r = ratio_of(m)
        vf = [f"crop=w='trunc(min(iw\\,ih*{r:.6f})/2)*2':h='trunc(min(ih\\,iw/{r:.6f})/2)*2'"]
    video_tool(eng, j, src, '_裁切', vf=vf + ['setsar=1'])

def t_rotate(eng, j, o):
    src = j.files[0]; op = o.get('op', 'cw')
    if o.get('fast') and op in ('cw', 'ccw', '180') and vext(src) in ('mp4', 'mov', 'mkv'):
        need_ff()
        out = out_for(j, src, vext(src), '_旋轉'); j.output = out
        deg = {'cw': -90, 'ccw': 90, '180': 180}[op]
        eng._proc(j, [FF(), '-hide_banner', '-nostdin', '-y', '-v', 'error', '-display_rotation:v:0', str(deg), '-i', src, '-map', '0', '-c', 'copy', out], '旋轉')
        j.encoder = '只改方向標記'; return
    vf = {'cw': ['transpose=1'], 'ccw': ['transpose=2'], '180': ['hflip', 'vflip'], 'hflip': ['hflip'], 'vflip': ['vflip']}[op]
    video_tool(eng, j, src, '_旋轉' if op in ('cw', 'ccw', '180') else '_翻轉', vf=vf)

def t_speed(eng, j, o):
    src = j.files[0]; f = max(0.1, min(16.0, num(o, 'speed', 2)))
    d = info(src)['dur'] / f
    keep = o.get('keep_audio', True) and has_audio(src)
    vf = [f'setpts=PTS/{f:.5f}']
    if keep:
        video_tool(eng, j, src, f'_{f:g}x', vf=vf, af=atempo_chain(f), dur=d)
    else:
        video_tool(eng, j, src, f'_{f:g}x', vf=vf, dur=d, audio=False)

def t_wm_text(eng, j, o):
    src = j.files[0]; font = find_font(bold=True)
    if not font: raise RuntimeError('找不到可用的字型')
    text = (o.get('text') or '').strip()
    if not text: raise RuntimeError('請輸入浮水印文字')
    tmp = tmpdir()
    try:
        with open(os.path.join(tmp, 'wm.txt'), 'w', encoding='utf-8') as fh: fh.write(text)
        size = num(o, 'size', 5) / 100; op = num(o, 'opacity', 80) / 100; col = hexcolor(o.get('color'))
        x, y = pos_xy(o.get('pos', 'br'), W='w', H='h', w='text_w', h='text_h', m='h*0.03')
        dt = (f"drawtext=fontfile='{ffesc(font)}':textfile=wm.txt:expansion=none:fontsize=h*{size:.4f}:fontcolor={col}@{op:.2f}:x={x}:y={y}"
              + (f":shadowcolor=black@{op * 0.6:.2f}:shadowx=2:shadowy=2:borderw=1:bordercolor=black@{op * 0.4:.2f}" if o.get('shadow', True) else ''))
        video_tool(eng, j, src, '_浮水印', vf=[dt], cwd=tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

def t_wm_img(eng, j, o):
    src = j.files[0]; logo = o.get('logo') or ''
    if not os.path.isfile(logo): raise RuntimeError('請選擇 Logo 圖片')
    W = info(src)['width'] or 1280
    lw = max(8, even(W * num(o, 'size', 15) / 100)); op = num(o, 'opacity', 90) / 100
    x, y = pos_xy(o.get('pos', 'br'), m='H*0.03')
    fc = f'[1:v]scale={lw}:-1,format=rgba,colorchannelmixer=aa={op:.3f}[lg];[0:v][lg]overlay={x}:{y}:format=auto,format=yuv420p[v]'
    video_tool(eng, j, src, '_Logo', fc=fc, maps=['[v]'] + (['0:a:0'] if has_audio(src) else []), extra_in=[['-i', logo]])
    # a still logo is a single frame: overlay repeats it (eof_action=repeat) until the main video ends

def t_delogo(eng, j, o):
    src = j.files[0]; x, y, w, h = parse_rect(o.get('rect'))
    i = info(src); W, H = i['width'] or 99999, i['height'] or 99999
    x = max(1, min(x, W - 3)); y = max(1, min(y, H - 3)); w = max(2, min(w, W - x - 1)); h = max(2, min(h, H - y - 1))
    video_tool(eng, j, src, '_去浮水印', vf=[f'delogo=x={x}:y={y}:w={w}:h={h}'])

def read_text_any(p):
    raw = open(p, 'rb').read()
    for enc in ('utf-8-sig', 'cp950', 'big5hkscs', 'gb18030', 'utf-16', 'cp1252'):
        try: return raw.decode(enc)
        except UnicodeDecodeError: continue
    return raw.decode('utf-8', 'replace')

def burn_subs(eng, j, src, sub, size=22, outline=True, suffix='_字幕', base=0.0, span=1.0):
    tmp = tmpdir()
    try:
        ext = os.path.splitext(sub)[1].lower() or '.srt'
        with open(os.path.join(tmp, 'sub' + ext), 'w', encoding='utf-8') as fh: fh.write(read_text_any(sub))
        style = f"FontName={font_family()},FontSize={int(size)},PrimaryColour=&H00FFFFFF,OutlineColour=&H00000000,BorderStyle=1,Outline={2 if outline else 0},Shadow={1 if outline else 0},MarginV=24"
        font_dir = os.path.dirname(find_font() or '') or tmp
        vf = [f"subtitles=sub{ext}:fontsdir='{ffesc(font_dir)}'" + (f":force_style='{style}'" if ext in ('.srt', '.vtt') else '')]
        return video_tool(eng, j, src, suffix, vf=vf, cwd=tmp, base=base, span=span)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

def soft_subs(eng, j, src, sub, suffix='_字幕'):
    need_ff()
    e = vext(src); e = 'mkv' if e == 'webm' else e
    out = out_for(j, src, e, suffix); j.output = out
    tmp = tmpdir()
    try:
        ext = os.path.splitext(sub)[1].lower() or '.srt'
        s2 = os.path.join(tmp, 'sub' + ext)
        with open(s2, 'w', encoding='utf-8') as fh: fh.write(read_text_any(sub))
        scodec = 'mov_text' if e in ('mp4', 'mov') else ('ass' if ext in ('.ass', '.ssa') else 'srt')
        eng._proc(j, [FF(), '-hide_banner', '-nostdin', '-y', '-v', 'error', '-i', src, '-sub_charenc', 'UTF-8', '-i', s2, '-map', '0:v', '-map', '0:a?', '-map', '1:0',
                      '-c', 'copy', '-c:s', scodec, '-metadata:s:s:0', 'language=chi', '-disposition:s:0', 'default', out], '加入字幕')
        j.encoder = '內嵌字幕（不重新編碼）'
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return out

def t_subs(eng, j, o):
    src = j.files[0]; sub = o.get('sub') or ''
    if not os.path.isfile(sub): raise RuntimeError('請選擇字幕檔')
    if o.get('mode', 'burn') == 'soft': soft_subs(eng, j, src, sub)
    else: burn_subs(eng, j, src, sub, num(o, 'size', 22), o.get('outline', True))

def t_audio(eng, j, o):
    src = j.files[0]; mus = o.get('music') or ''
    if not os.path.isfile(mus): raise RuntimeError('請選擇音樂檔')
    need_ff()
    d = info(src)['dur']; mv = num(o, 'music_vol', 35) / 100; ov = num(o, 'orig_vol', 100) / 100; fd = num(o, 'fade', 2)
    mix = o.get('mode', 'mix') == 'mix' and has_audio(src)
    fade = f',afade=t=out:st={max(0, d - fd):.3f}:d={fd:.3f}' if fd > 0 and d > fd else ''
    if mix:
        fc = f'[1:a]volume={mv:.3f}{fade}[m];[0:a]volume={ov:.3f}[o];[o][m]amix=inputs=2:duration=first:dropout_transition=0:normalize=0[a]'
    else:
        fc = f'[1:a]volume={mv:.3f}{fade},apad[a]'
    out = out_for(j, src, vext(src), '_配樂'); j.output = out
    loop = ['-stream_loop', '-1'] if o.get('loop', True) else []
    t = os.path.splitext(out)[1].lstrip('.')
    c = base_cmd() + ['-i', src] + loop + ['-i', mus, '-filter_complex', fc, '-map', '0:v:0', '-map', '[a]', '-c:v', 'copy',
                                           '-c:a', 'libopus' if t == 'webm' else 'aac', '-b:a', f"{E.PRESET[o.get('preset', 'balanced')]['aac']}k", '-t', f'{d:.3f}']
    c += (['-movflags', '+faststart'] if t in ('mp4', 'mov') else []) + [out]
    j.encoder = 'FFmpeg（畫面不重新編碼）'
    eng._exec(j, c, d)
    if j._rc != 0 and not j.cancelled: raise RuntimeError(ff_error(j._err, '配樂失敗'))

def t_mute(eng, j, o):
    need_ff(); src = j.files[0]
    out = out_for(j, src, vext(src), '_靜音'); j.output = out
    eng._proc(j, [FF(), '-hide_banner', '-nostdin', '-y', '-v', 'error', '-i', src, '-map', '0:v', '-c', 'copy', '-an'] +
              (['-movflags', '+faststart'] if out.endswith(('.mp4', '.mov')) else []) + [out], '去除聲音')
    j.encoder = '不重新編碼'

def t_split(eng, j, o):
    need_ff(); src = j.files[0]; d = info(src)['dur']
    if d <= 0: raise RuntimeError('讀不到影片長度')
    seg = num(o, 'sec', 60) if o.get('mode', 'every') == 'every' else d / max(2, int(num(o, 'parts', 3)))
    seg = max(0.5, seg)
    if seg >= d: raise RuntimeError(f'每段 {seg:g} 秒，比整個檔案（{d:.1f} 秒）還長')
    folder = out_dir_for(j, src, '_分割'); j.output = folder
    isv = has_video(src)
    ext = vext(src) if isv else (os.path.splitext(src)[1].lstrip('.').lower() or 'mp3')
    fmt = ['-segment_format', 'ipod'] if ext == 'm4a' else []
    pat = os.path.join(folder, os.path.splitext(os.path.basename(src))[0] + f'_%03d.{ext}')
    expected = math.ceil(d / seg - 0.02)

    def go(exact):
        for f in os.listdir(folder): os.remove(os.path.join(folder, f))
        c = base_cmd() + ['-i', src] + (['-map', '0:v:0'] if isv else []) + ['-map', '0:a?']
        if exact and isv:
            c += venc_args(j, pat.replace('%03d', '000'), d, E.detect_hw() if o.get('hwaccel', True) else [])
            c += ['-force_key_frames', f'expr:gte(t,n_forced*{seg:.3f})']
        elif exact:
            c += E.audio_args(ext, o) if (ext in E.AUDIO_COPY or ext == 'wav') else ['-c:a', 'aac', '-b:a', '192k']
            j.encoder = 'FFmpeg（精準切割）'
        else:
            c += ['-c', 'copy']; j.encoder = '不重新編碼（在關鍵影格切）'
        c += ['-f', 'segment', '-segment_time', f'{seg:.3f}', '-segment_time_delta', '0.05', '-reset_timestamps', '1'] + fmt + [pat]
        eng._exec(j, c, d)
        if j._rc != 0 and not j.cancelled: raise RuntimeError(ff_error(j._err, '分割失敗'))
        return len(os.listdir(folder))

    exact = bool(o.get('exact')) or not isv          # audio re-encode is fast and always exact
    n = go(exact)
    if not exact and n < expected * 0.75:
        n = go(True); j.encoder += '（原檔關鍵影格太少，已自動改用精準切割）'
    j.encoder += f'，共 {n} 段'

def t_stab(eng, j, o):
    need_ff(); src = j.files[0]; d = info(src)['dur']; sm = int(num(o, 'strength', 15))
    tmp = tmpdir()
    try:
        c1 = base_cmd() + ['-i', src, '-vf', 'vidstabdetect=shakiness=6:accuracy=9:stepsize=6:result=t.trf', '-an', '-f', 'null', '-']
        j.encoder = '分析晃動中…'
        eng._exec(j, c1, d, cwd=tmp, base=0, span=0.45)
        if j._rc != 0 and not j.cancelled: raise RuntimeError('分析晃動失敗：' + (j._err.strip().splitlines() or [''])[-1][:200])
        video_tool(eng, j, src, '_防手震', vf=[f'vidstabtransform=input=t.trf:smoothing={sm}:zoom=0:optzoom=1:interpol=bicubic', 'unsharp=5:5:0.6:3:3:0.3'],
                   cwd=tmp, base=0.45, span=0.55)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

def t_enhance(eng, j, o):
    src = j.files[0]; vf = []
    b = num(o, 'bright', 0) / 400; c = 1 + num(o, 'contrast', 0) / 100; s = 1 + num(o, 'sat', 0) / 100
    if b or c != 1 or s != 1: vf.append(f'eq=brightness={b:.3f}:contrast={c:.3f}:saturation={s:.3f}')
    dn = o.get('denoise', '0')
    if dn == '1': vf.append('hqdn3d=2:1.5:3:2.25')
    elif dn == '2': vf.append('hqdn3d=5:4:7:5')
    if o.get('sharpen'): vf.append('unsharp=5:5:0.8:3:3:0.0')
    if not vf: raise RuntimeError('沒有要調整的項目')
    video_tool(eng, j, src, '_調整', vf=vf)

def t_fade(eng, j, o):
    src = j.files[0]; d = info(src)['dur']; a = num(o, 'fin', 1); b = num(o, 'fout', 1)
    vf, af = [], []
    if a > 0: vf.append(f'fade=t=in:st=0:d={a}'); af.append(f'afade=t=in:st=0:d={a}')
    if b > 0 and d > b: vf.append(f'fade=t=out:st={d - b:.3f}:d={b}'); af.append(f'afade=t=out:st={d - b:.3f}:d={b}')
    if not vf: raise RuntimeError('淡入、淡出秒數都是 0')
    video_tool(eng, j, src, '_淡入淡出', vf=vf, af=af)

def t_reverse(eng, j, o):
    src = j.files[0]; d = info(src)['dur']
    if d > 300: raise RuntimeError('影片超過 5 分鐘，倒轉會用掉大量記憶體，請先用「分割」切短')
    video_tool(eng, j, src, '_倒轉', vf=['reverse'], af=['areverse'] if o.get('audio', True) else None, audio=bool(o.get('audio', True)))

def t_frames(eng, j, o):
    need_ff(); src = j.files[0]; i = info(src); d = i['dur']
    if o.get('mode', 'sheet') == 'every':
        folder = out_dir_for(j, src, '_畫面'); j.output = folder
        sec = max(0.04, num(o, 'sec', 5))
        c = base_cmd() + ['-i', src, '-vf', ','.join(E.hdr_filters(src) + [f'fps=1/{sec:.4f}']), '-q:v', '3', os.path.join(folder, '%04d.jpg')]
        j.encoder = 'FFmpeg'; eng._exec(j, c, d)
    else:
        cols = int(num(o, 'cols', 4)); rows = int(num(o, 'rows', 5)); n = cols * rows
        if d <= 0: raise RuntimeError('讀不到影片長度')
        out = out_for(j, src, 'jpg', '_縮圖總覽'); j.output = out
        font = find_font()
        stamp = (f",drawtext=fontfile='{ffesc(font)}':text='%{{pts\\:hms}}':x=8:y=h-th-8:fontsize=22:fontcolor=white:box=1:boxcolor=black@0.5:boxborderw=4") if font else ''
        vf = ''.join(x + ',' for x in E.hdr_filters(src)) + f"fps={n / d:.6f},scale=480:-2{stamp},tile={cols}x{rows}:padding=6:margin=6:color=white"
        c = base_cmd() + ['-i', src, '-vf', vf, '-frames:v', '1', '-q:v', '3', '-update', '1', out]
        j.encoder = 'FFmpeg'; eng._exec(j, c, d)
    if j._rc != 0 and not j.cancelled: raise RuntimeError(ff_error(j._err, '擷取失敗'))

def t_grid(eng, j, o):
    fs = j.files
    if not 2 <= len(fs) <= 4: raise RuntimeError('請選 2～4 支影片')
    lay = o.get('layout', 'h'); H = int(o.get('size') or 720)
    if lay == 'grid' and len(fs) != 4: raise RuntimeError('田字排列需要剛好 4 支影片')
    n = len(fs); d = min(info(f)['dur'] for f in fs)
    if lay == 'grid':
        W = even(H * 16 / 9)
        pre = ''.join(f'[{i}:v]scale={W}:{H}:force_original_aspect_ratio=decrease,pad={W}:{H}:(ow-iw)/2:(oh-ih)/2,setsar=1,fps=30[v{i}];' for i in range(4))
        fc = pre + '[v0][v1][v2][v3]xstack=inputs=4:layout=0_0|w0_0|0_h0|w0_h0:shortest=1,format=yuv420p[v]'
    elif lay == 'h':
        pre = ''.join(f'[{i}:v]scale=-2:{H},setsar=1,fps=30[v{i}];' for i in range(n))
        fc = pre + ''.join(f'[v{i}]' for i in range(n)) + f'hstack=inputs={n}:shortest=1,format=yuv420p[v]'
    else:
        Wd = even(H * 16 / 9)
        pre = ''.join(f'[{i}:v]scale={Wd}:-2,setsar=1,fps=30[v{i}];' for i in range(n))
        fc = pre + ''.join(f'[v{i}]' for i in range(n)) + f'vstack=inputs={n}:shortest=1,format=yuv420p[v]'
    for i in range(1, n):
        tm = E.hdr_filters(fs[i])
        if tm: fc = fc.replace(f'[{i}:v]', f'[{i}:v]' + ','.join(tm) + ',', 1)
    maps = ['[v]'] + (['0:a:0'] if has_audio(fs[0]) else [])
    video_tool(eng, j, fs[0], '_並排', fc=fc, maps=maps, extra_in=[['-i', f] for f in fs[1:]], dur=d, ext='mp4')

# ---------------------------------------------------------------- AI subtitles (whisper.cpp inside ffmpeg)
MODELS = {'base': ('ggml-base.bin', 147951465), 'small': ('ggml-small.bin', 487601967), 'medium': ('ggml-medium.bin', 1533763059),
          'turbo': ('ggml-large-v3-turbo-q5_0.bin', 574041195)}
MODEL_URL = 'https://huggingface.co/ggerganov/whisper.cpp/resolve/main/'

def model_dir():
    """Models the user downloads later (writable)."""
    d = os.path.join(E.data_dir(), 'models'); os.makedirs(d, exist_ok=True); return d

def find_model(name, min_size=0):
    """Bundled models (tools\\models, shipped with the app) first, then ones downloaded later."""
    for d in (E.bundled('models'), model_dir()):
        p = os.path.join(d, name)
        if os.path.isfile(p) and os.path.getsize(p) >= min_size: return p
    return None

_WHISPER = None
def whisper_ok():
    """Does the bundled ffmpeg have the whisper filter?  Asked while the window opens, so the answer is kept
    on disk per ffmpeg build: starting the 220 MB ffmpeg.exe just to list filters cost up to a second per launch."""
    global _WHISPER
    if _WHISPER is None:
        if not E.FFMPEG: _WHISPER = False; return _WHISPER
        cache = os.path.join(E.data_dir(), 'ffmpeg_caps.json')
        try: sig = [os.path.getmtime(E.FFMPEG), os.path.getsize(E.FFMPEG)]
        except OSError: sig = None
        try:
            with open(cache, encoding='utf-8') as fh: c = json.load(fh)
            if sig and c.get('ffmpeg') == sig and isinstance(c.get('whisper'), bool): _WHISPER = c['whisper']; return _WHISPER
        except Exception: pass
        try: _WHISPER = ' whisper ' in E.run([E.FFMPEG, '-hide_banner', '-filters']).stdout
        except Exception: _WHISPER = False; return _WHISPER
        try:
            with open(cache, 'w', encoding='utf-8') as fh: json.dump(dict(ffmpeg=sig, whisper=_WHISPER), fh)
        except Exception: pass
    return _WHISPER

_MODEL_LOCK = threading.Lock()
def ensure_model(j, key):
    with _MODEL_LOCK:
        if j.cancelled: raise InterruptedError
        return _ensure_model(j, key)

def _ensure_model(j, key):
    name, size = MODELS.get(key) or MODELS['base']
    have = find_model(name, size * 0.98)
    if have: return have
    for k in installed_models():                     # the chosen model is gone: use any bundled one
        return find_model(MODELS[k][0], MODELS[k][1] * 0.98)
    raise RuntimeError('缺少離線語音辨識模型（tools\\models\\ggml-base.bin），請重新安裝完整版本')

VAD_URL = 'https://huggingface.co/ggml-org/whisper-vad/resolve/main/ggml-silero-v5.1.2.bin'
def ensure_vad():
    with _MODEL_LOCK: return _ensure_vad()

def _ensure_vad():
    have = find_model('ggml-silero-v5.1.2.bin', 100000)
    if have: return have
    return None

def to_trad(text):
    try:
        import zhconv; return zhconv.s2twp(text)
    except Exception:
        return text

def whisper_srt(eng, j, src, o, base=0.0, span=1.0):
    """Offline speech recognition (whisper.cpp inside ffmpeg) -> cleaned SRT text (Traditional Chinese if asked)."""
    need_ff()
    if not whisper_ok(): raise RuntimeError('這個 FFmpeg 版本沒有內建語音辨識（需要 FFmpeg 8 以上的 full 版）')
    if not has_audio(src): raise RuntimeError('檔案裡沒有聲音')
    if o.get('model') == 'medium' and E.mem_gb() < 6: raise RuntimeError('記憶體較少，請使用「快速」或「準確」模型，避免系統過度使用分頁檔')
    model = ensure_model(j, o.get('model', 'base'))
    lang = o.get('lang', 'zh'); lang = 'zh' if lang == 'yue' else lang   # whisper.cpp has no separate Cantonese in older models
    d = info(src)['dur']
    tmp = tmpdir()
    try:
        vad = ensure_vad()
        maxlen = 22 if lang in ('zh', 'ja', 'ko', 'yue') else 42
        # VAD cuts the audio at real pauses -> better sentence boundaries and no "hallucinated" text in silence
        variants = []
        if vad: variants.append(f":queue=20:vad_model='{ffesc(vad)}':max_len={maxlen}")
        variants += [f':queue=10:max_len={maxlen}', ':queue=10']
        j.encoder = 'AI 語音辨識中…'
        for k, extra in enumerate(variants):
            af = f"aresample=16000,whisper=model='{ffesc(model)}':language={lang}{extra}:destination=out.srt:format=srt"
            c = base_cmd() + ['-i', src, '-vn', '-af', af, '-f', 'null', '-']
            eng._exec(j, c, d, cwd=tmp, base=base, span=span)
            if j._rc == 0 or j.cancelled: break
            if k == len(variants) - 1 or not re.search(r'Option .* not found|vad|Error applying option', j._err, re.I):
                raise RuntimeError('語音辨識失敗：' + ff_error(j._err))
        if j.cancelled: raise InterruptedError
        sp = os.path.join(tmp, 'out.srt')
        if not os.path.isfile(sp) or os.path.getsize(sp) < 5: raise RuntimeError('沒有辨識到任何語音')
        with open(sp, encoding='utf-8', errors='replace') as fh: text = fix_srt(fh.read())
        if o.get('trad', True) and lang in ('zh', 'auto'): text = to_trad(text)
        return text
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

def t_autosub(eng, j, o):
    src = j.files[0]
    base = 0.3 if j.progress >= 0.29 else 0.0
    burn = o.get('out', 'srt') == 'burn'
    span = (0.5 if burn else 1.0) * (1 - base)
    text = whisper_srt(eng, j, src, o, base, span)
    srt = out_for(j, src, 'srt', '')
    with open(srt, 'w', encoding='utf-8-sig') as fh: fh.write(text)
    j.output = srt; j.keep_output = True
    if o.get('out') == 'burn' and has_video(src):
        burn_subs(eng, j, src, srt, 22, True, '_AI字幕', base=base + span, span=1 - base - span)
    elif o.get('out') == 'soft' and has_video(src):
        soft_subs(eng, j, src, srt, '_AI字幕')
    if o.get('out') != 'srt': j.encoder = 'AI 字幕（另存 SRT：' + os.path.basename(srt) + '）'
    else: j.encoder = 'AI 語音辨識（whisper）'

# ---------------------------------------------------------------- transcripts & meeting notes
def srt_cues(text):
    """SRT text -> [(start_seconds, end_seconds, text)]"""
    def sec(t):
        h, m, rest = t.strip().replace('.', ',').split(':'); s, _, ms = rest.partition(',')
        return int(h) * 3600 + int(m) * 60 + int(s) + int((ms or '0')[:3].ljust(3, '0')) / 1000
    cues = []
    for b in re.split(r'\n\s*\n', text.replace('\r', '').strip()):
        lines = b.strip().split('\n')
        k = next((i for i, l in enumerate(lines) if '-->' in l), None)
        if k is None: continue
        a, _, z = lines[k].partition('-->')
        try: st, en = sec(a), sec(z.split()[0])
        except (ValueError, IndexError): continue
        body = ' '.join(l.strip() for l in lines[k + 1:] if l.strip())
        if body: cues.append((st, en, body))
    return cues

def hms(t):
    t = int(t); return f'{t // 3600:02d}:{t % 3600 // 60:02d}:{t % 60:02d}'

def paragraphs(cues, cjk=True, gap=2.0, limit=260):
    """Join recognised lines into readable paragraphs: a new paragraph after a pause or when one gets long."""
    out = []; cur = []; start = None; last = None
    joiner = '' if cjk else ' '
    for st, en, tx in cues:
        if cur and (st - last > gap or len(joiner.join(cur)) > limit):
            out.append((start, joiner.join(cur))); cur = []
        if not cur: start = st
        if cjk and cur and not re.search(r'[，。！？、,.!?；;：:]$', cur[-1]): cur[-1] += '，'
        cur.append(tx); last = en
    if cur: out.append((start, joiner.join(cur)))
    return [(s, re.sub(r'，$', '', p) + ('。' if cjk and not re.search(r'[。！？.!?]$', p) else '')) for s, p in out]

def write_docx(path, title, blocks):
    """blocks: [('h', text) | ('p', text) | ('b', text) bullet | ('t', (time, text)) timestamped paragraph]"""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    doc = Document()
    font = font_family() if find_font() else 'Microsoft JhengHei'
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet'):
        try:
            st = doc.styles[name]; st.font.name = font
            st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), font)
        except KeyError: pass
    doc.styles['Normal'].font.size = Pt(11)
    doc.add_heading(title, 0)
    for kind, val in blocks:
        if kind == 'h': doc.add_heading(val, 1)
        elif kind == 'b': doc.add_paragraph(val, style='List Bullet')
        elif kind == 't':
            p = doc.add_paragraph(); r = p.add_run(f'[{hms(val[0])}] '); r.font.color.rgb = RGBColor(0x86, 0x86, 0x8B); p.add_run(val[1])
        else: doc.add_paragraph(val)
    doc.save(path)

def t_transcribe(eng, j, o):
    src = j.files[0]; fmt = o.get('fmt', 'docx')
    text = whisper_srt(eng, j, src, o, 0.0, 0.95)
    cjk = o.get('lang', 'zh') in ('zh', 'ja', 'ko', 'yue', 'auto')
    paras = paragraphs(srt_cues(text), cjk)
    if not paras: raise RuntimeError('沒有辨識到任何語音')
    stem = os.path.splitext(os.path.basename(src))[0]
    if fmt == 'srt':
        out = out_for(j, src, 'srt', '_逐字稿')
        with open(out, 'w', encoding='utf-8-sig') as fh: fh.write(text)
    elif fmt == 'docx':
        out = out_for(j, src, 'docx', '_逐字稿')
        write_docx(out, stem + ' 逐字稿', [('p', f'來源：{os.path.basename(src)}　長度：{hms(info(src)["dur"])}　（離線語音辨識，請核對專有名詞）')] + [('t', p) for p in paras])
    else:
        out = out_for(j, src, 'txt', '_逐字稿')
        with open(out, 'w', encoding='utf-8-sig') as fh:
            fh.write('\n\n'.join((f'[{hms(s)}] ' if fmt == 'txt_ts' else '') + p for s, p in paras) + '\n')
    j.output = out; j.encoder = f'AI 語音辨識（whisper，{len(paras)} 段）'

SUMMARY_SECTIONS = ('摘要', '重點', '決議', '待辦事項')

def _sections(text):
    """'【摘要】…【重點】…' or '摘要：' / '## 重點' headings -> {section: [lines]}"""
    out = {k: [] for k in SUMMARY_SECTIONS}; cur = '重點'
    for line in text.splitlines():
        l = line.strip().strip('*').strip()
        if not l: continue
        hit = next((k for k in SUMMARY_SECTIONS if re.fullmatch(rf'[#\s【\[]*{k}[】\]：:\s]*', l) or re.match(rf'^[#\s]*[【\[]{k}[】\]]', l)), None)
        if hit:
            cur = hit; rest = re.sub(rf'^[#\s]*[【\[]?{hit}[】\]]?[：:]?\s*', '', l)
            if rest: out[cur].append(rest)
            continue
        l = re.sub(r'^([-•・*]|\d+[.、)])\s*', '', l)
        if l and not re.fullmatch(r'(無|沒有|N/?A|none)[。.]?', l, re.I): out[cur].append(l)
    return out

def t_meeting(eng, j, o):
    src = j.files[0]; mode = o.get('mode', 'ai'); n = max(3, min(20, int(num(o, 'points', 8))))
    if mode == 'ai' and not ai_ok(): raise RuntimeError('缺少離線文字模型（tools\\pdf-ai），請改用「快速摘錄」或重新安裝完整版本')
    text = whisper_srt(eng, j, src, o, 0.0, 0.55 if mode == 'ai' else 0.9)
    cjk = o.get('lang', 'zh') in ('zh', 'ja', 'ko', 'yue', 'auto')
    paras = paragraphs(srt_cues(text), cjk)
    if not paras: raise RuntimeError('沒有辨識到任何語音')
    transcript = '\n'.join(p for _, p in paras)
    import pdf_local_engines as L
    if mode == 'ai':
        lang_name = '繁體中文' if cjk else '與原文相同的語言'
        pieces = L.chunks(transcript, 1000); notes = []
        for i, part in enumerate(pieces):
            if j.cancelled: raise InterruptedError
            j.encoder = f'AI 整理會議內容 {i + 1}/{len(pieces)}'
            notes.append(L.generate(j, f'以下是會議逐字稿片段（語音辨識，可能有錯字）。請用{lang_name}條列這段的重點，最多 5 點；'
                                       '有做出決定的，用「決議：」開頭；有人要負責的事，用「待辦：」開頭並寫出負責人與期限（逐字稿有提到才寫）。只輸出條列，不要編造。', part, 320))
            j.progress = 0.55 + 0.3 * (i + 1) / len(pieces)
        merged = '\n'.join(notes)
        while len(merged) > 1400:                        # several hours of meeting: condense the notes first
            parts = L.chunks(merged, 1300)
            merged = '\n'.join(L.generate(j, f'把以下會議筆記濃縮成最多 6 點{lang_name}條列，保留「決議：」和「待辦：」項目。', p, 320) for p in parts)
        j.encoder = 'AI 產生會議總結'
        final = L.generate(j, f'根據以下會議筆記，用{lang_name}寫出會議總結，嚴格使用這四個標題：\n【摘要】2～3 句話說明會議主題與結論\n【重點】最多 {n} 點條列\n'
                              '【決議】條列，沒有就寫「無」\n【待辦事項】條列「負責人：事項（期限）」，沒有就寫「無」\n只根據筆記內容，不要編造。', merged, 700)
        sec = _sections(final)
        for line in merged.splitlines():                 # decisions / to-dos the final pass dropped
            l = re.sub(r'^([-•・*]|\d+[.、)])\s*', '', line.strip())
            for key, tag in (('決議', '決議：'), ('待辦事項', '待辦：')):
                if l.startswith(tag) and not any(l[len(tag):].strip()[:8] in x for x in sec[key]): sec[key].append(l[len(tag):].strip())
        if not sec['摘要'] and sec['重點']: sec['摘要'] = [sec['重點'][0]]
        engine = 'whisper + llama.cpp（離線）'
    else:
        sec = {k: [] for k in SUMMARY_SECTIONS}
        sec['重點'] = [l.lstrip('• ') for l in L.extractive(transcript, n).splitlines() if l.strip()]
        sec['摘要'] = sec['重點'][:1]
        for s, p in paras:
            for sentence in re.split(r'(?<=[。！？.!?])', p):
                if re.search(r'決定|決議|確定|通過|同意|定案', sentence): sec['決議'].append(sentence.strip())
                elif re.search(r'負責|處理|完成|跟進|追蹤|下週|下次|之前|截止|due|deadline|follow', sentence, re.I): sec['待辦事項'].append(sentence.strip())
        sec['決議'] = list(dict.fromkeys(x for x in sec['決議'] if x))[:10]; sec['待辦事項'] = list(dict.fromkeys(x for x in sec['待辦事項'] if x))[:10]
        engine = 'whisper + 快速摘錄（離線）'
    j.progress = 0.9
    title = (o.get('title') or '').strip() or os.path.splitext(os.path.basename(src))[0]
    folder = out_dir_for(j, src, '_會議總結'); j.output = folder
    meta = f'來源：{os.path.basename(src)}　長度：{hms(info(src)["dur"])}　產生方式：{engine}'
    blocks = [('p', meta)]
    for k in SUMMARY_SECTIONS:
        blocks.append(('h', k))
        items = sec[k] or ['無']
        blocks += [('p' if k == '摘要' else 'b', x) for x in items]
    blocks += [('h', '逐字稿')] + [('t', p) for p in paras]
    write_docx(os.path.join(folder, title + '_會議總結.docx'), title + ' 會議總結', blocks)
    with open(os.path.join(folder, title + '_會議總結.md'), 'w', encoding='utf-8') as fh:
        fh.write(f'# {title} 會議總結\n\n{meta}\n\n' + '\n\n'.join(f'## {k}\n\n' + '\n'.join(('' if k == '摘要' else '- ') + x for x in (sec[k] or ['無'])) for k in SUMMARY_SECTIONS)
                 + '\n\n## 逐字稿\n\n' + '\n\n'.join(f'[{hms(s)}] {p}' for s, p in paras) + '\n')
    with open(os.path.join(folder, '逐字稿.srt'), 'w', encoding='utf-8-sig') as fh: fh.write(text)
    j.encoder = f'會議總結（{engine}）'
    j.notes = ['語音辨識與 AI 整理可能有錯字或遺漏，重要決議與待辦請對照逐字稿確認。']

def fix_srt(text):
    """Renumber cues and drop empty ones (whisper sometimes emits blanks)."""
    blocks = re.split(r'\n\s*\n', text.replace('\r', '').strip())
    out = []
    for b in blocks:
        lines = b.strip().split('\n')
        k = next((i for i, l in enumerate(lines) if '-->' in l), None)
        if k is None: continue
        body = '\n'.join(l.strip() for l in lines[k + 1:] if l.strip())
        if not body: continue
        out.append(f'{len(out) + 1}\n{lines[k].strip()}\n{body}')
    return '\n\n'.join(out) + '\n'

# ---------------------------------------------------------------- audio tools
def t_volume(eng, j, o):
    src = j.files[0]; d = info(src)['dur']; af = []
    g = num(o, 'gain', 0); a = num(o, 'fin', 0); b = num(o, 'fout', 0)
    if o.get('norm'): af.append('loudnorm=I=-16:TP=-1.5:LRA=11')
    if g: af.append(f'volume={g}dB')
    if g > 0: af.append('alimiter=limit=0.97:level=false')
    if a > 0: af.append(f'afade=t=in:st=0:d={a}')
    if b > 0 and d > b: af.append(f'afade=t=out:st={d - b:.3f}:d={b}')
    if not af: raise RuntimeError('沒有要調整的項目')
    audio_tool(eng, j, src, '_音量', af)

def t_pitch(eng, j, o):
    src = j.files[0]; t = num(o, 'tempo', 100) / 100; s = num(o, 'semi', 0)
    if t == 1 and s == 0: raise RuntimeError('速度 100%、音調 0，沒有變化')
    af = [f'rubberband=tempo={t:.4f}:pitch={2 ** (s / 12):.6f}:pitchq=quality:formant=preserved' if s else f'rubberband=tempo={t:.4f}:pitchq=quality']
    suffix = (f'_{s:+g}key' if s else '') + (f'_{t * 100:g}%' if t != 1 else '')
    if has_video(src) and o.get('fmt', 'same') == 'same':
        video_tool(eng, j, src, suffix, vf=[f'setpts=PTS/{t:.5f}'], af=af, dur=info(src)['dur'] / t)
    else:
        audio_tool(eng, j, src, suffix, af, dur=info(src)['dur'] / t)

def t_vocal(eng, j, o):
    src = j.files[0]
    if 'mono' in E.run([FF(), '-hide_banner', '-i', src]).stderr.split('Audio:', 1)[-1].split('\n', 1)[0]:
        raise RuntimeError('這是單聲道音訊，無法去人聲（需要立體聲）')
    if o.get('bass', True):
        fc = '[0:a]asplit=2[x][y];[x]pan=stereo|c0=c0-c1|c1=c1-c0[s];[y]lowpass=f=160,pan=stereo|c0=0.5*c0+0.5*c1|c1=0.5*c0+0.5*c1[l];[s][l]amix=inputs=2:normalize=0[a]'
    else:
        fc = '[0:a]pan=stereo|c0=c0-c1|c1=c1-c0[a]'
    audio_tool(eng, j, src, '_伴奏', None, fc=fc, maps=['[a]'])

def t_denoise(eng, j, o):
    src = j.files[0]; nr = int(num(o, 'level', 12))
    audio_tool(eng, j, src, '_降噪', ['highpass=f=70', f'afftdn=nr={nr}:nf=-40:tn=1', 'lowpass=f=15000'])

def t_silence(eng, j, o):
    src = j.files[0]; db = num(o, 'db', -45); m = num(o, 'min', 0.8)
    af = [f'silenceremove=start_periods=1:start_threshold={db}dB:stop_periods=-1:stop_duration={m}:stop_threshold={db}dB']
    audio_tool(eng, j, src, '_去靜音', af)

def t_mix(eng, j, o):
    need_ff(); fs = j.files
    if len(fs) < 2: raise RuntimeError('混音至少要選 2 個檔案')
    n = len(fs); dmode = o.get('dur', 'longest')
    ds = [info(f)['dur'] for f in fs]; d = {'longest': max(ds), 'first': ds[0], 'shortest': min(ds)}[dmode]
    fc = ''.join(f'[{i}:a:0]' for i in range(n)) + f'amix=inputs={n}:duration={dmode}:dropout_transition=0:normalize=0,alimiter=limit=0.97:level=false[a]'
    ext = o.get('fmt', 'mp3')
    out = out_for(j, fs[0], ext, '_混音'); j.output = out
    c = base_cmd()
    for f in fs: c += ['-i', f]
    c += ['-filter_complex', fc, '-map', '[a]'] + E.audio_args(ext, o) + [out]
    j.encoder = 'FFmpeg'; eng._exec(j, c, d)
    if j._rc != 0 and not j.cancelled: raise RuntimeError(ff_error(j._err, '混音失敗'))

def t_ringtone(eng, j, o):
    need_ff(); src = j.files[0]; st = num(o, 'start', 0); ln = num(o, 'len', 30); fmt = o.get('fmt', 'm4r')
    if fmt == 'm4r': ln = min(ln, 40)
    out = out_for(j, src, fmt, '_鈴聲'); j.output = out
    af = f'afade=t=in:st=0:d=0.5,afade=t=out:st={max(0, ln - 1.5):.2f}:d=1.5'
    c = base_cmd(['-ss', str(st)]) + ['-t', str(ln), '-i', src, '-vn', '-map', '0:a:0', '-af', af]
    c += (['-c:a', 'aac', '-b:a', '192k', '-f', 'ipod'] if fmt == 'm4r' else ['-c:a', 'libmp3lame', '-b:a', '192k']) + [out]
    j.encoder = 'FFmpeg'; eng._exec(j, c, ln)
    if j._rc != 0 and not j.cancelled: raise RuntimeError(ff_error(j._err, '鈴聲製作失敗'))

def t_tags(eng, j, o):
    need_ff(); src = j.files[0]; ext = os.path.splitext(src)[1].lstrip('.').lower()
    out = out_for(j, src, ext, '_資訊'); j.output = out
    cover = o.get('cover') or ''
    c = [FF(), '-hide_banner', '-nostdin', '-y', '-v', 'error', '-i', src]
    if cover and os.path.isfile(cover) and ext in ('mp3', 'm4a', 'flac'):
        c += ['-i', cover, '-map', '0:a', '-map', '1:0', '-c', 'copy', '-disposition:v:0', 'attached_pic']
        if ext == 'mp3': c += ['-metadata:s:v', 'title=Album cover', '-metadata:s:v', 'comment=Cover (front)']
    else:
        c += ['-map', '0', '-c', 'copy']
    for k, key in (('title', 'title'), ('artist', 'artist'), ('album', 'album'), ('year', 'date')):
        v = (o.get(k) or '').strip()
        if v: c += ['-metadata', f'{key}={v}']
    if ext == 'mp3': c += ['-id3v2_version', '3']
    eng._proc(j, c + [out], '寫入歌曲資訊'); j.encoder = '不重新編碼'

# ---------------------------------------------------------------- image tools
def open_img(path, want=None):
    from PIL import ImageOps
    return ImageOps.exif_transpose(E.Engine._open_img(None, path, want))

def img_ext(src, o):
    f = o.get('fmt') or 'same'
    if f != 'same': return f
    e = os.path.splitext(src)[1].lower().lstrip('.')
    return {'jpeg': 'jpg', 'jfif': 'jpg', 'tif': 'tiff'}.get(e, e) if e in ('jpg', 'jpeg', 'jfif', 'png', 'webp', 'bmp', 'tif', 'tiff') else 'jpg'

def save_img(im, out, o, icc=None):
    t = os.path.splitext(out)[1].lstrip('.').lower(); q = E.PRESET[o.get('preset', 'balanced')]['q']
    kw = {}
    if t == 'jpg': im = E.flatten(im); kw = dict(format='JPEG', quality=q, optimize=True, progressive=True)
    elif t == 'png': kw = dict(format='PNG', optimize=True)
    elif t == 'webp': kw = dict(format='WEBP', quality=q, method=5)
    elif t == 'bmp': im = E.flatten(im); kw = dict(format='BMP')
    elif t == 'tiff': kw = dict(format='TIFF', compression='tiff_adobe_deflate')
    if icc and t in ('jpg', 'png', 'webp', 'tiff'): kw['icc_profile'] = icc
    im.save(out, **kw)

def t_resize(eng, j, o):
    from PIL import Image
    src = j.files[0]; m = o.get('mode', 'long'); v = num(o, 'val', 1920)
    factor = lambda W, H: {'long': v / max(W, H), 'w': v / W, 'h': v / H, 'pct': v / 100}[m]
    W0 = H0 = None
    def want(W, H):                                    # remember the true size: JPEG may be decoded pre-shrunk
        nonlocal W0, H0; W0, H0 = W, H; return factor(W, H)
    im = open_img(src, want); icc = im.info.get('icc_profile')
    W, H = (W0, H0) if W0 else im.size
    s = factor(W, H)
    if s < 1: im = im.resize((max(1, round(W * s)), max(1, round(H * s))), Image.LANCZOS, reducing_gap=3.0)
    out = out_for(j, src, img_ext(src, o), '_縮放'); j.output = out
    save_img(im, out, o, icc); j.encoder = f'Pillow（{im.width}×{im.height}）'

def crop_ratio(im, r):
    W, H = im.size
    if W / H > r: w = round(H * r); x = (W - w) // 2; return im.crop((x, 0, x + w, H))
    h = round(W / r); y = (H - h) // 2; return im.crop((0, y, W, y + h))

def t_irotate(eng, j, o):
    from PIL import Image
    src = j.files[0]; im = open_img(src); icc = im.info.get('icc_profile')
    op = o.get('op', 'none')
    tr = {'cw': Image.Transpose.ROTATE_270, 'ccw': Image.Transpose.ROTATE_90, '180': Image.Transpose.ROTATE_180,
          'hflip': Image.Transpose.FLIP_LEFT_RIGHT, 'vflip': Image.Transpose.FLIP_TOP_BOTTOM}.get(op)
    cr = o.get('crop', 'none')
    if cr == 'custom':
        x, y, w, h = parse_rect(o.get('rect')); im = im.crop((x, y, x + w, y + h))   # rect is picked on the un-rotated image
    if tr is not None: im = im.transpose(tr)
    if cr not in ('none', 'custom'): im = crop_ratio(im, ratio_of(cr))
    if tr is None and cr == 'none': raise RuntimeError('沒有選擇要旋轉或裁切')
    out = out_for(j, src, img_ext(src, o), '_編輯'); j.output = out
    save_img(im, out, o, icc); j.encoder = 'Pillow'

def t_iwm(eng, j, o):
    from PIL import Image, ImageDraw, ImageFont
    src = j.files[0]; base = open_img(src); icc = base.info.get('icc_profile')
    base = base.convert('RGBA'); W, H = base.size
    op = num(o, 'opacity', 60) / 100; size = num(o, 'size', 20) / 100
    if o.get('kind', 'text') == 'logo':
        lp = o.get('logo') or ''
        if not os.path.isfile(lp): raise RuntimeError('請選擇 Logo 圖片')
        mark = Image.open(lp).convert('RGBA'); w = max(4, round(W * size)); mark = mark.resize((w, max(1, round(mark.height * w / mark.width))), Image.LANCZOS)
    else:
        text = (o.get('text') or '').strip()
        if not text: raise RuntimeError('請輸入浮水印文字')
        fp = find_font(bold=True)
        target_w = W * size; fs = 40
        font = ImageFont.truetype(fp, fs) if fp else ImageFont.load_default()
        bw = ImageDraw.Draw(Image.new('L', (1, 1))).textbbox((0, 0), text, font=font)[2] or 1
        if fp: fs = max(8, int(fs * target_w / bw)); font = ImageFont.truetype(fp, fs)
        bb = ImageDraw.Draw(Image.new('L', (1, 1))).textbbox((0, 0), text, font=font, stroke_width=max(1, fs // 40))
        mark = Image.new('RGBA', (bb[2] - bb[0] + 4, bb[3] - bb[1] + 4), (0, 0, 0, 0))
        col = hexcolor(o.get('color')); rgb = tuple(int(col[i:i + 2], 16) for i in (1, 3, 5))
        ImageDraw.Draw(mark).text((2 - bb[0], 2 - bb[1]), text, font=font, fill=rgb + (255,), stroke_width=max(1, fs // 40), stroke_fill=(0, 0, 0, 90))
    a = mark.split()[3].point(lambda v: int(v * op)); mark.putalpha(a)
    layer = Image.new('RGBA', base.size, (0, 0, 0, 0))
    pos = o.get('pos', 'br')
    if pos == 'tile':
        rot = mark.rotate(30, expand=True, resample=Image.BICUBIC)
        sx, sy = int(rot.width * 1.6) + 1, int(rot.height * 2.2) + 1
        for yy, y in enumerate(range(-rot.height, H + rot.height, sy)):
            for x in range(-rot.width + (yy % 2) * sx // 2, W + rot.width, sx):
                layer.paste(rot, (x, y), rot)
    else:
        m = round(min(W, H) * 0.03); mw, mh = mark.size
        xs = {'l': m, 'c': (W - mw) // 2, 'r': W - mw - m}; ys = {'t': m, 'c': (H - mh) // 2, 'b': H - mh - m}
        x, y = {'br': ('r', 'b'), 'bl': ('l', 'b'), 'tr': ('r', 't'), 'tl': ('l', 't'), 'c': ('c', 'c'), 'b': ('c', 'b'), 't': ('c', 't')}[pos]
        layer.paste(mark, (xs[x], ys[y]), mark)
    im = Image.alpha_composite(base, layer)
    out = out_for(j, src, img_ext(src, o), '_浮水印'); j.output = out
    save_img(im, out, o, icc); j.encoder = 'Pillow'

def t_collage(eng, j, o):
    from PIL import Image
    fs = j.files
    if len(fs) < 2: raise RuntimeError('至少要選 2 張圖片')
    lay = o.get('layout', 'v'); S = int(num(o, 'size', 1080)); gap = int(num(o, 'gap', 0))
    col = hexcolor(o.get('bg')); bg = tuple(int(col[i:i + 2], 16) for i in (1, 3, 5))
    ims = []
    for i, f in enumerate(fs):
        if j.cancelled: raise InterruptedError
        ims.append(E.flatten(open_img(f, lambda W, H: S / min(W, H) if lay == 'grid' else S / (W if lay == 'v' else H)), bg)); j.progress = 0.8 * (i + 1) / len(fs)
    if lay == 'v':
        ims = [im.resize((S, max(1, round(im.height * S / im.width))), Image.LANCZOS, reducing_gap=3.0) for im in ims]
        canvas = Image.new('RGB', (S, sum(im.height for im in ims) + gap * (len(ims) - 1)), bg); y = 0
        for im in ims: canvas.paste(im, (0, y)); y += im.height + gap
    elif lay == 'h':
        ims = [im.resize((max(1, round(im.width * S / im.height)), S), Image.LANCZOS, reducing_gap=3.0) for im in ims]
        canvas = Image.new('RGB', (sum(im.width for im in ims) + gap * (len(ims) - 1), S), bg); x = 0
        for im in ims: canvas.paste(im, (x, 0)); x += im.width + gap
    else:
        cols = int(num(o, 'cols', 3)); rows = math.ceil(len(ims) / cols); cell = (S - gap * (cols - 1)) // cols
        canvas = Image.new('RGB', (cols * cell + gap * (cols - 1), rows * cell + gap * (rows - 1)), bg)
        for k, im in enumerate(ims):
            im = crop_ratio(im, 1.0).resize((cell, cell), Image.LANCZOS, reducing_gap=3.0)
            canvas.paste(im, ((k % cols) * (cell + gap), (k // cols) * (cell + gap)))
    if max(canvas.size) > 65000: raise RuntimeError('拼出來的圖太大（超過 65000 像素），請減少張數或寬度')
    ext = o.get('fmt', 'jpg')
    if ext == 'jpg' and max(canvas.size) > 65500: ext = 'png'
    out = out_for(j, fs[0], ext, '_拼圖'); j.output = out
    save_img(canvas, out, o); j.encoder = f'Pillow（{canvas.width}×{canvas.height}）'

def t_slideshow(eng, j, o):
    need_ff(); fs = j.files
    W, H = [int(v) for v in o.get('size', '1920x1080').split('x')]
    D = max(0.5, num(o, 'sec', 3)); tr = o.get('trans', 'fade'); T = min(num(o, 'tdur', 0.6), D * 0.45)
    if tr == 'none' or len(fs) == 1: T = 0
    tmp = tmpdir()
    try:
        # normalise every picture to a JPEG first (HEIC / RAW / PSD / transparent PNG all work the same way)
        from PIL import Image
        paths = []
        for i, f in enumerate(fs):
            if j.cancelled: raise InterruptedError
            im = E.flatten(open_img(f, lambda w, h: min(W * 2 / w, H * 2 / h)), (0, 0, 0)); im.thumbnail((W * 2, H * 2), Image.LANCZOS)
            p = os.path.join(tmp, f'{i:04d}.jpg'); im.save(p, quality=92); paths.append(p)
            j.progress = 0.1 * (i + 1) / len(fs)
        n = len(paths); total = n * D - (n - 1) * T
        c = base_cmd()
        for p in paths: c += ['-loop', '1', '-framerate', '30', '-t', f'{D:.3f}', '-i', p]
        mus = o.get('music') or ''
        if os.path.isfile(mus): c += ['-stream_loop', '-1', '-i', mus]
        pre = ''.join(f'[{i}:v]split=2[a{i}][b{i}];[a{i}]scale={W}:{H}:force_original_aspect_ratio=increase,crop={W}:{H},boxblur=24:2,eq=brightness=-0.08[g{i}];'
                      f'[b{i}]scale={W}:{H}:force_original_aspect_ratio=decrease[f{i}];[g{i}][f{i}]overlay=(W-w)/2:(H-h)/2,setsar=1,fps=30,format=yuv420p[v{i}];' for i in range(n))
        if n == 1: chain = '[v0]null[v]'
        elif T > 0:
            chain = ''; prev = 'v0'
            for i in range(1, n):
                nxt = 'v' if i == n - 1 else f'x{i}'
                chain += f'[{prev}][v{i}]xfade=transition={tr}:duration={T:.3f}:offset={i * (D - T):.3f}[{nxt}];'; prev = nxt
            chain = chain.rstrip(';')
        else:
            chain = ''.join(f'[v{i}]' for i in range(n)) + f'concat=n={n}:v=1:a=0[v]'
        fc = pre + chain
        if os.path.isfile(mus): fc += f';[{n}:a]afade=t=out:st={max(0, total - 2):.3f}:d=2[a]'
        out = out_for(j, fs[0], 'mp4', '_影片'); j.output = out
        def build(hw):
            cc = c + ['-filter_complex', fc, '-map', '[v]'] + (['-map', '[a]'] if os.path.isfile(mus) else [])
            va = venc_args(j, out, total, hw, audio=os.path.isfile(mus))
            return cc + va + ['-t', f'{total:.3f}', out]
        ff_run(eng, j, build, total, base=0.1, span=0.9)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

def t_gif(eng, j, o):
    from PIL import Image
    fs = j.files; w = int(num(o, 'w', 600)); frames = []
    for i, f in enumerate(fs):
        if j.cancelled: raise InterruptedError
        im = E.flatten(open_img(f, lambda W, H: w / W)); im = im.resize((w, max(1, round(im.height * w / im.width))), Image.LANCZOS, reducing_gap=3.0)
        frames.append(im); j.progress = 0.7 * (i + 1) / len(fs)
    H = max(f.height for f in frames)
    frames = [f if f.height == H else (lambda c: (c.paste(f, (0, (H - f.height) // 2)), c)[1])(Image.new('RGB', (w, H), (255, 255, 255))) for f in frames]
    pal = [f.convert('P', palette=Image.ADAPTIVE, colors=256, dither=Image.FLOYDSTEINBERG) for f in frames]
    out = out_for(j, fs[0], 'gif', '_動畫'); j.output = out
    pal[0].save(out, save_all=True, append_images=pal[1:], duration=int(num(o, 'ms', 500)), loop=0 if o.get('loop', True) else 1, optimize=True, disposal=2)
    j.encoder = 'Pillow'

# ---------------------------------------------------------------- PDF tools
def parse_pages(spec, n):
    """'1, 3-5, 8-' -> [0, 2, 3, 4, 7, ...] (keeps the order given, 1-based in, 0-based out)."""
    out = []
    for part in re.split(r'[,，、;\s]+', (spec or '').strip()):
        if not part: continue
        m = re.fullmatch(r'(\d*)\s*[-~－–]\s*(\d*)', part)
        if m:
            a = int(m.group(1) or 1); b = int(m.group(2) or n)
            rng = range(a, b + 1) if a <= b else range(a, b - 1, -1)
            out += [p - 1 for p in rng if 1 <= p <= n]
        elif part.isdigit() and 1 <= int(part) <= n:
            out.append(int(part) - 1)
        else:
            raise RuntimeError(f'看不懂的頁碼：{part}')
    if not out: raise RuntimeError('沒有符合的頁碼')
    return out

def pdf_reader(path, pw=None):
    from pypdf import PdfReader
    r = PdfReader(path)
    if r.is_encrypted:
        if not pw or not r.decrypt(pw): raise RuntimeError('這份 PDF 有密碼，請先用「PDF 解除密碼」')
    return r

def write_pdf(w, out):
    try: w.compress_identical_objects(remove_identicals=True, remove_orphans=True)
    except Exception: pass
    with open(out, 'wb') as fh: w.write(fh)

def t_psplit(eng, j, o):
    from pypdf import PdfWriter
    src = j.files[0]; r = pdf_reader(src); n = len(r.pages)
    m = o.get('mode', 'each')
    if m == 'each': groups = [[i] for i in range(n)]
    elif m == 'every':
        k = max(1, int(num(o, 'n', 5))); groups = [list(range(i, min(n, i + k))) for i in range(0, n, k)]
    else:
        groups = [parse_pages(p, n) for p in re.split(r'[,，、;]+', o.get('ranges') or '') if p.strip()]
    folder = out_dir_for(j, src, '_分割'); j.output = folder
    stem = os.path.splitext(os.path.basename(src))[0]
    for gi, g in enumerate(groups):
        if j.cancelled: raise InterruptedError
        w = PdfWriter()
        for p in g: w.add_page(r.pages[p])
        name = f'{stem}_p{g[0] + 1}' + (f'-{g[-1] + 1}' if len(g) > 1 else '') + '.pdf'
        write_pdf(w, os.path.join(folder, name)); j.progress = (gi + 1) / len(groups)
    j.encoder = f'pypdf（{len(groups)} 個檔案）'

def t_ppages(eng, j, o):
    from pypdf import PdfWriter
    src = j.files[0]; r = pdf_reader(src); n = len(r.pages)
    sel = parse_pages(o.get('pages'), n)
    keep = sel if o.get('mode', 'keep') == 'keep' else [i for i in range(n) if i not in set(sel)]
    if not keep: raise RuntimeError('刪除後就沒有頁面了')
    w = PdfWriter()
    for p in keep: w.add_page(r.pages[p])
    out = out_for(j, src, 'pdf', '_整理'); j.output = out
    write_pdf(w, out); j.encoder = f'pypdf（{len(keep)} 頁）'

def t_protate(eng, j, o):
    from pypdf import PdfWriter
    src = j.files[0]; r = pdf_reader(src); n = len(r.pages)
    sel = set(parse_pages(o.get('pages'), n)) if (o.get('pages') or '').strip() else set(range(n))
    w = PdfWriter(clone_from=r); a = int(o.get('angle', '90'))
    for i, pg in enumerate(w.pages):
        if i in sel: pg.rotate(a)
    out = out_for(j, src, 'pdf', '_旋轉'); j.output = out
    write_pdf(w, out); j.encoder = 'pypdf'

def t_pencrypt(eng, j, o):
    from pypdf import PdfWriter
    from pypdf.constants import UserAccessPermissions as P
    src = j.files[0]; pw = o.get('pw') or ''
    if not pw: raise RuntimeError('請輸入密碼')
    r = pdf_reader(src); w = PdfWriter(clone_from=r)
    perm = P(0)
    if o.get('print', True): perm |= P.PRINT | P.PRINT_TO_REPRESENTATION
    if o.get('copy'): perm |= P.EXTRACT | P.EXTRACT_TEXT_AND_GRAPHICS
    import secrets
    owner = secrets.token_hex(16)
    try: w.encrypt(user_password=pw, owner_password=owner, algorithm='AES-256', permissions_flag=perm); alg = 'AES-256'
    except ImportError: raise RuntimeError('缺少 AES 加密元件，請使用包含 cryptography 的完整建置版本')
    out = out_for(j, src, 'pdf', '_加密'); j.output = out
    with open(out, 'wb') as fh: w.write(fh)
    j.encoder = f'pypdf（{alg}）'

def t_pdecrypt(eng, j, o):
    from pypdf import PdfReader, PdfWriter
    src = j.files[0]; r = PdfReader(src)
    if not r.is_encrypted: raise RuntimeError('這份 PDF 沒有密碼')
    try: ok = r.decrypt(o.get('pw') or '')
    except Exception as e: raise RuntimeError('無法解密：' + str(e)[:120])
    if not ok: raise RuntimeError('密碼錯誤')
    w = PdfWriter(clone_from=r)
    out = out_for(j, src, 'pdf', '_無密碼'); j.output = out
    write_pdf(w, out); j.encoder = 'pypdf'

# ---------------------------------------------------------------- privacy
def strip_jpeg(data):
    """Drop APP1 (EXIF/XMP), APP12/13 (IPTC/Photoshop) and COM without touching the image data."""
    if data[:2] != b'\xff\xd8': return None
    out = bytearray(b'\xff\xd8'); i = 2
    while i + 4 <= len(data):
        if data[i] != 0xFF: return None
        mk = data[i + 1]
        if mk == 0xDA: out += data[i:]; return bytes(out)          # start of scan: copy the rest
        if mk in (0x01,) or 0xD0 <= mk <= 0xD7: out += data[i:i + 2]; i += 2; continue
        ln = int.from_bytes(data[i + 2:i + 4], 'big')
        seg = data[i:i + 2 + ln]
        if not (mk == 0xE1 or mk in (0xEC, 0xED) or mk == 0xFE): out += seg
        i += 2 + ln
    return None

def strip_png(data):
    if data[:8] != b'\x89PNG\r\n\x1a\n': return None
    out = bytearray(data[:8]); i = 8
    while i + 8 <= len(data):
        ln = int.from_bytes(data[i:i + 4], 'big'); typ = data[i + 4:i + 8]
        chunk = data[i:i + 12 + ln]
        if typ not in (b'eXIf', b'tEXt', b'iTXt', b'zTXt', b'tIME'): out += chunk
        i += 12 + ln
        if typ == b'IEND': break
    return bytes(out)

def t_privacy(eng, j, o):
    src = j.files[0]; cat, ext = E.classify(src)
    out = out_for(j, src, ext, '_隱私清除'); j.output = out
    if cat == 'image':
        data = open(src, 'rb').read()
        res = strip_jpeg(data) if ext in ('jpg', 'jpeg', 'jfif') else strip_png(data) if ext == 'png' else None
        if res:
            with open(out, 'wb') as fh: fh.write(res)
            j.encoder = '無損移除 EXIF／GPS'; return
        im = open_img(src); icc = im.info.get('icc_profile')
        if ext in ('heic', 'heif') or ext in E.RAW or ext in ('psd', 'svg'):
            out2 = out[:-len(ext)] + 'jpg'; j.output = out2; save_img(im, out2, o, icc)
        else:
            t = img_ext(src, {}); save_img(im, out, dict(o, preset='high'), icc) if t in ('jpg', 'png', 'webp', 'tiff', 'bmp') else im.save(out)
        j.encoder = 'Pillow（重新存檔，不含 EXIF）'; return
    if cat == 'pdf':
        from pypdf import PdfWriter
        r = pdf_reader(src); w = PdfWriter(clone_from=r)
        try: w.metadata = None
        except Exception: w.add_metadata({'/Producer': '', '/Creator': '', '/Author': '', '/Title': ''})
        try:
            if '/Metadata' in w._root_object: del w._root_object['/Metadata']
        except Exception: pass
        write_pdf(w, out); j.encoder = 'pypdf（移除作者與 XMP）'; return
    need_ff()
    eng._proc(j, [FF(), '-hide_banner', '-nostdin', '-y', '-v', 'error', '-i', src, '-map', '0', '-map_metadata', '-1', '-map_chapters', '-1',
                  '-c', 'copy', '-fflags', '+bitexact', '-flags:v', '+bitexact', '-flags:a', '+bitexact'] +
              (['-movflags', '+faststart'] if ext in ('mp4', 'mov', 'm4a') else []) + [out], '清除資訊')
    j.encoder = '移除拍攝時間、地點等資訊（不重新編碼）'

# ---------------------------------------------------------------- disc tools
def list_drives():
    """Optical drives (Windows) as [{'path': 'D:\\', 'label': ...}]."""
    out = []
    if os.name != 'nt': return out
    import ctypes, string
    k = ctypes.windll.kernel32
    mask = k.GetLogicalDrives()
    for i, L in enumerate(string.ascii_uppercase):
        if mask >> i & 1:
            root = f'{L}:\\'
            if k.GetDriveTypeW(root) == 5:
                name = ctypes.create_unicode_buffer(261)
                ok = k.GetVolumeInformationW(root, name, 261, None, None, None, None, 0)
                out.append(dict(path=root, label=f'{L}: ' + (name.value if ok and name.value else ('（沒有光碟）' if not ok else '光碟'))))
    return out

def find_video_ts(path):
    for p in (path, os.path.join(path, 'VIDEO_TS')):
        if os.path.isdir(p) and any(f.upper().endswith('.VOB') for f in os.listdir(p)): return p
    if os.path.isfile(path) and path.upper().endswith('.VOB'): return os.path.dirname(path)
    return None

def t_dvd(eng, j, o):
    need_ff(); src = j.files[0]
    if src.lower().endswith('.iso'): raise RuntimeError('ISO 請先用 7-Zip 解開（壓縮檔分頁可直接解壓），再選 VIDEO_TS 資料夾')
    vt = find_video_ts(src)
    if not vt: raise RuntimeError('找不到 VIDEO_TS（請放入 DVD 或選擇 VIDEO_TS 資料夾）')
    groups = {}
    for f in os.listdir(vt):
        m = re.fullmatch(r'VTS_(\d\d)_([1-9])\.VOB', f.upper())
        if m: groups.setdefault(m.group(1), []).append(os.path.join(vt, f))
    if not groups: raise RuntimeError('VIDEO_TS 裡沒有影片')
    title = max(groups.values(), key=lambda fs: sum(os.path.getsize(f) for f in fs))
    title.sort()
    ext = o.get('fmt', 'mp4'); label = os.path.basename(os.path.dirname(vt.rstrip('\\/'))) or 'DVD'
    out = out_for(j, os.path.join(os.path.dirname(vt), label), ext, '_DVD'); j.output = out
    inp = 'concat:' + '|'.join(title)
    pre = ['-analyzeduration', '100M', '-probesize', '100M', '-fflags', '+genpts']
    d = 0.0
    r = E.run([FF(), '-hide_banner'] + pre + ['-i', inp])
    m = re.search(r'Duration: (\d+):(\d+):([\d.]+)', r.stderr)
    if m: d = int(m.group(1)) * 3600 + int(m.group(2)) * 60 + float(m.group(3))
    if not d: d = sum(os.path.getsize(f) for f in title) / (5_000_000 / 8)   # ~5 Mbps estimate
    vf = (['bwdif=mode=send_field:deint=interlaced'] if o.get('deint', True) else []) + ['setsar=1']
    def build(hw):
        c = base_cmd(pre) + ['-i', inp, '-vf', ','.join(vf)]
        c += (['-map', '0:v:0', '-map', '0:a?', '-map', '0:s?', '-c:s', 'dvdsub'] if ext == 'mkv' else ['-map', '0:v:0', '-map', '0:a:0?'])
        return c + venc_args(j, out, d, hw) + [out]
    try:
        ff_run(eng, j, build, d)
    except RuntimeError as e:
        if 'Invalid data' in str(e) or 'decrypt' in str(e).lower():
            raise RuntimeError('這片 DVD 有防拷保護（CSS），無法轉換')
        raise

def t_bluray(eng, j, o):
    need_ff(); src = j.files[0].rstrip('\\/') + os.sep
    if not os.path.isdir(os.path.join(src, 'BDMV')): raise RuntimeError('找不到 BDMV 資料夾（請放入藍光光碟或選擇光碟根目錄）')
    ext = o.get('fmt', 'mp4'); out = out_for(j, src.rstrip('\\/'), ext, '_藍光'); j.output = out
    inp = 'bluray:' + src
    r = E.run([FF(), '-hide_banner', '-i', inp])
    if 'aacs' in r.stderr.lower() or 'encrypted' in r.stderr.lower(): raise RuntimeError('這片藍光有防拷保護（AACS），無法轉換')
    m = re.search(r'Duration: (\d+):(\d+):([\d.]+)', r.stderr)
    d = int(m.group(1)) * 3600 + int(m.group(2)) * 60 + float(m.group(3)) if m else 0
    def build(hw):
        c = base_cmd() + ['-i', inp, '-map', '0:v:0'] + (['-map', '0:a?'] if ext == 'mkv' else ['-map', '0:a:0?'])
        return c + venc_args(j, out, d, hw) + [out]
    ff_run(eng, j, build, d)

def cd_toc(drive):
    import ctypes
    from ctypes import wintypes
    k = ctypes.windll.kernel32
    if not isinstance(drive,str) or not re.fullmatch(r'[A-Za-z]:\\?',drive):raise RuntimeError('請選擇音樂光碟機')
    if k.GetDriveTypeW(str(drive[0])+':\\') != 5:raise RuntimeError('請選擇音樂光碟機')
    k.CreateFileW.argtypes=[ctypes.c_wchar_p,wintypes.DWORD,wintypes.DWORD,ctypes.c_void_p,wintypes.DWORD,wintypes.DWORD,wintypes.HANDLE]
    k.CreateFileW.restype=wintypes.HANDLE
    k.DeviceIoControl.argtypes=[wintypes.HANDLE,wintypes.DWORD,ctypes.c_void_p,wintypes.DWORD,ctypes.c_void_p,wintypes.DWORD,ctypes.POINTER(wintypes.DWORD),ctypes.c_void_p]
    k.DeviceIoControl.restype=wintypes.BOOL
    k.CloseHandle.argtypes=[wintypes.HANDLE];k.CloseHandle.restype=wintypes.BOOL
    h = k.CreateFileW(f'\\\\.\\{drive[0]}:', 0x80000000, 3, None, 3, 0, None)
    if h in (-1, 0xFFFFFFFF, ctypes.c_void_p(-1).value): raise RuntimeError('無法開啟光碟機')
    buf = (ctypes.c_ubyte * 804)(); ret = wintypes.DWORD()
    if not k.DeviceIoControl(ctypes.c_void_p(h), 0x24000, None, 0, buf, 804, ctypes.byref(ret), None):
        k.CloseHandle(ctypes.c_void_p(h)); raise RuntimeError('讀不到光碟目錄（請確認放入的是音樂 CD）')
    first, last = buf[2], buf[3]; tracks = []
    for t in range(first, last + 2):
        o = 4 + (t - first) * 8
        ctrl = buf[o + 1] & 0x0F                      # TRACK_DATA: Control is the low nibble; bit 2 = data track
        m_, s_, f_ = buf[o + 5], buf[o + 6], buf[o + 7]
        tracks.append((buf[o + 2], (m_ * 60 + s_) * 75 + f_ - 150, ctrl))
    return h, tracks

def t_cd(eng, j, o):
    if os.name != 'nt': raise RuntimeError('音樂 CD 只支援 Windows')
    import ctypes, wave
    from ctypes import wintypes
    need_ff(); k = ctypes.windll.kernel32
    drive = j.files[0]; h, toc = cd_toc(drive)
    try:
        audio = [(toc[i][0], toc[i][1], toc[i + 1][1]) for i in range(len(toc) - 1) if not (toc[i][2] & 0x4)]
        if not audio: raise RuntimeError('這片光碟沒有音樂音軌')
        folder = j.opt.get('out_dir') or os.path.join(default_dir(j), '音樂CD_' + time.strftime('%Y%m%d_%H%M'))
        folder = E.unique_path(folder); os.makedirs(folder); j.output = folder
        fmt = o.get('fmt', 'flac'); total = sum(b - a for _, a, b in audio); done = 0;unreadable=0
        class RAW(ctypes.Structure): _fields_ = [('DiskOffset', ctypes.c_longlong), ('SectorCount', wintypes.ULONG), ('TrackMode', ctypes.c_int)]
        N = 20; buf = (ctypes.c_ubyte * (2352 * N))(); ret = wintypes.DWORD()
        for no, a, b in audio:
            wav = os.path.join(folder, f'Track{no:02d}.wav')
            with wave.open(wav, 'wb') as w:
                w.setnchannels(2); w.setsampwidth(2); w.setframerate(44100)
                lba = a
                while lba < b:
                    if j.cancelled: raise InterruptedError
                    n = min(N, b - lba); ri = RAW(lba * 2048, n, 2)
                    ok = False
                    for _ in range(3):
                        if k.DeviceIoControl(ctypes.c_void_p(h), 0x2403E, ctypes.byref(ri), ctypes.sizeof(ri), buf, 2352 * n, ctypes.byref(ret), None): ok = True; break
                    if not ok:unreadable+=n
                    w.writeframes(bytes(buf[:ret.value]) if ok else b'\0' * (2352 * n))
                    lba += n; done += n; j.progress = 0.85 * done / total
            if fmt != 'wav':
                outp = wav[:-3] + fmt
                r = E.run([FF(), '-hide_banner', '-nostdin', '-y', '-v', 'error', '-i', wav] + E.audio_args(fmt, o) + ['-metadata', f'track={no}', outp])
                if r.returncode == 0: os.remove(wav)
                else:j.warnings.append(f'第 {no} 軌音訊轉換未完成，已保留 WAV 原始擷取檔')
        if unreadable:j.warnings.append(f'{unreadable} 個光碟區塊無法讀取，已以靜音填補；請核對音軌')
        j.encoder = f'{len(audio)} 首音軌'
    finally:
        k.CloseHandle(ctypes.c_void_p(h))

def t_iso(eng, j, o):
    import ctypes
    src = j.files[0]
    if os.name != 'nt' or not re.fullmatch(r'[A-Za-z]:\\?', src): raise RuntimeError('請選擇光碟機')
    if ctypes.windll.kernel32.GetDriveTypeW(str(src[0])+':\\') != 5:raise RuntimeError('請選擇光碟機')
    dev = f'\\\\.\\{src[0]}:'
    import ctypes
    free = ctypes.c_ulonglong(); total = ctypes.c_ulonglong()
    ctypes.windll.kernel32.GetDiskFreeSpaceExW(f'{src[0]}:\\', None, ctypes.byref(total), ctypes.byref(free))
    size = total.value or 1
    name = ctypes.create_unicode_buffer(261)
    ctypes.windll.kernel32.GetVolumeInformationW(f'{src[0]}:\\', name, 261, None, None, None, None, 0)
    out = out_for(j, os.path.join(default_dir(j), name.value or 'disc'), 'iso', ''); j.output = out
    got = 0; t0 = time.time()
    with open(dev, 'rb', buffering=0) as fi, open(out, 'wb') as fo:
        while True:
            if j.cancelled: raise InterruptedError
            try: b = fi.read(2048 * 512)
            except OSError: break
            if not b: break
            fo.write(b); got += len(b)
            j.progress = min(0.99, got / size); el = time.time() - t0
            j.speed = f'{got / 1048576 / max(el, 0.1):.1f} MB/s'
    if got == 0: raise RuntimeError('讀不到光碟內容（音樂 CD 請用「音樂 CD 轉檔」）')
    j.encoder = f'ISO（{got / 1048576:.0f} MB）'

RUN = dict(v_social=t_social, v_fill=t_fill, v_crop=t_crop, v_rotate=t_rotate, v_speed=t_speed, v_wm_text=t_wm_text, v_wm_img=t_wm_img,
           v_delogo=t_delogo, v_subs=t_subs, v_autosub=t_autosub, v_audio=t_audio, v_mute=t_mute, v_split=t_split, v_stab=t_stab,
           v_enhance=t_enhance, v_fade=t_fade, v_reverse=t_reverse, v_frames=t_frames, v_grid=t_grid,
           a_volume=t_volume, a_pitch=t_pitch, a_vocal=t_vocal, a_denoise=t_denoise, a_silence=t_silence, a_mix=t_mix, a_ringtone=t_ringtone, a_tags=t_tags,
           i_resize=t_resize, i_rotate=t_irotate, i_watermark=t_iwm, i_collage=t_collage, i_slideshow=t_slideshow, i_gif=t_gif,
           p_split=t_psplit, p_pages=t_ppages, p_rotate=t_protate, p_encrypt=t_pencrypt, p_decrypt=t_pdecrypt,
           x_privacy=t_privacy, d_dvd=t_dvd, d_bluray=t_bluray, d_cd=t_cd, d_iso=t_iso, a_transcribe=t_transcribe, a_meeting=t_meeting)
def _offline_pdf(eng,j,o):
    import pdf_tools
    return pdf_tools.run(eng,j,o)
RUN.update({t['id']:_offline_pdf for t in PDF_SCHEMA.TOOLS})
assert set(RUN) == set(TOOL), set(TOOL) ^ set(RUN)

# which tools mostly run ffmpeg (go to the small "heavy" pool)
HEAVY = {t['id'] for t in TOOLS if t['group'] in ('video', 'disc')} | {'a_pitch', 'a_transcribe', 'a_meeting'}
